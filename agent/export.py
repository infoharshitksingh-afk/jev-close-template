"""
Export the question log to Excel: one row per question, a summary by area (formulas), the
decisions owners made, and the full trail of who said what.
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .qlog import FIELDS

HEAD = PatternFill("solid", fgColor="1F3A5F")
WHITE = Font(color="FFFFFF", bold=True)
GREEN = PatternFill("solid", fgColor="D9F2E3")
RED = PatternFill("solid", fgColor="F9D9D9")
AMBER = PatternFill("solid", fgColor="FFF1CC")
WRAP = Alignment(wrap_text=True, vertical="top")

WIDTH = {"id": 9, "opened_at": 17, "close_month": 9, "asked_by": 20, "channel": 14, "thread_link": 18,
         "question": 45, "area": 16, "problem": 35, "agent_answer": 55, "evidence": 40, "status": 15,
         "escalated_to": 26, "escalation_reason": 20, "escalation_question": 45, "owner_answer": 45,
         "answered_by": 18, "answered_at": 17, "decision": 40, "change_type": 17, "worked": 9,
         "worked_note": 30, "closed_at": 17, "related_to": 10}

README = [
    ("Close agent question log", None),
    ("Every question the close agent was asked: the problem behind it, what it answered, who it asked "
     "when the data couldn't settle it, what they said, and whether it worked.", None),
    ("", None),
    ("Tab", "What it holds"),
    ("Log", "One row per question. Filter by area, status or Worked?. Green = solved, red = didn't work, amber = waiting on an owner."),
    ("Summary", "Counts by area (formulas on the Log tab): how many questions, how many the agent answered alone, how many needed an owner, how many worked."),
    ("Decisions", "Owner answers that settle something. Read these before the next close: they are the standing rules."),
    ("Trail", "Every event in order: asked, answered, escalated, owner answered, reminder, outcome."),
    ("", None),
    ("Status", "Meaning"),
    ("answered", "The agent answered from the data. Waiting for the asker to confirm it worked."),
    ("waiting_on_owner", "The agent tagged an owner and is waiting for their reply."),
    ("owner_answered", "The owner replied. Waiting for the asker to confirm it worked."),
    ("closed", "The asker confirmed it solved the problem."),
    ("", None),
    ("Why escalated", "Meaning"),
    ("judgment_call", "A policy or accounting choice the data can't make (LOGIC_REVIEW.md assumptions)."),
    ("sources_disagree", "Two sources give different numbers and the data can't say which is right."),
    ("contradicts_prior_decision", "The question or today's data conflicts with an owner's earlier answer."),
    ("unexplained_flag", "A check flagged and the fix needs a person."),
    ("would_change_numbers", "The answer means changing config, an input cell, a source file or an invoice."),
    ("data_not_available", "The close doesn't have the data to answer."),
]


def export(log, path):
    rows = log.all()
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    for i, (a, b) in enumerate(README, 1):
        rd.cell(i, 1, a)
        if b:
            rd.cell(i, 2, b).alignment = WRAP
        if b and a in ("Tab", "Status", "Why escalated"):
            rd.cell(i, 1).font = rd.cell(i, 2).font = Font(bold=True)
    rd["A1"].font = Font(bold=True, size=14)
    rd.column_dimensions["A"].width = 26
    rd.column_dimensions["B"].width = 110

    # ---- Log
    ws = wb.create_sheet("Log")
    for j, (k, label) in enumerate(FIELDS, 1):
        c = ws.cell(1, j, label)
        c.fill, c.font, c.alignment = HEAD, WHITE, WRAP
        ws.column_dimensions[get_column_letter(j)].width = WIDTH.get(k, 15)
    for i, r in enumerate(rows, 2):
        for j, (k, _) in enumerate(FIELDS, 1):
            c = ws.cell(i, j, r[k] or None)
            c.alignment = WRAP
            if k == "thread_link" and str(r[k] or "").startswith("http"):
                c.hyperlink, c.value, c.font = r[k], "open thread", Font(color="0563C1", underline="single")
    last = max(len(rows) + 1, 2)
    ref = f"A1:{get_column_letter(len(FIELDS))}{last}"
    ws.add_table(Table(displayName="QuestionLog", ref=ref,
                       tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=False)))
    col = {k: get_column_letter(j) for j, (k, _) in enumerate(FIELDS, 1)}
    rng = f"A2:{get_column_letter(len(FIELDS))}{last}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col["worked"]}2="yes"'], fill=GREEN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col["worked"]}2="no"'], fill=RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${col["status"]}2="waiting_on_owner"'], fill=AMBER))
    ws.freeze_panes = "C2"

    # ---- Summary (formulas)
    sm = wb.create_sheet("Summary", 1)
    heads = ["Area", "Questions", "Agent answered alone", "Needed an owner", "Waiting on owner",
             "Worked", "Didn't work", "Pending", "Solved rate"]
    for j, h in enumerate(heads, 1):
        c = sm.cell(1, j, h)
        c.fill, c.font, c.alignment = HEAD, WHITE, WRAP
        sm.column_dimensions[get_column_letter(j)].width = 16
    sm.column_dimensions["A"].width = 24
    areas = sorted({r["area"] or "(none)" for r in rows}) or ["(none)"]
    L = lambda k: f"Log!${col[k]}$2:${col[k]}${last}"  # noqa: E731  (exact rows: "" would count blank rows)
    for i, a in enumerate(areas, 2):
        key = "" if a == "(none)" else a
        sm.cell(i, 1, a)
        sm.cell(i, 2, f'=COUNTIFS({L("area")},"{key}")')
        sm.cell(i, 3, f'=COUNTIFS({L("area")},"{key}",{L("escalated_to")},"")')
        sm.cell(i, 4, f"=B{i}-C{i}")
        sm.cell(i, 5, f'=COUNTIFS({L("area")},"{key}",{L("status")},"waiting_on_owner")')
        sm.cell(i, 6, f'=COUNTIFS({L("area")},"{key}",{L("worked")},"yes")')
        sm.cell(i, 7, f'=COUNTIFS({L("area")},"{key}",{L("worked")},"no")')
        sm.cell(i, 8, f'=COUNTIFS({L("area")},"{key}",{L("worked")},"pending")')
        sm.cell(i, 9, f'=IF(F{i}+G{i}=0,"",F{i}/(F{i}+G{i}))').number_format = "0%"
    t = len(areas) + 2
    sm.cell(t, 1, "Total").font = Font(bold=True)
    for j in range(2, 9):
        L_ = get_column_letter(j)
        sm.cell(t, j, f"=SUM({L_}2:{L_}{t - 1})").font = Font(bold=True)
    sm.cell(t, 9, f'=IF(F{t}+G{t}=0,"",F{t}/(F{t}+G{t}))').number_format = "0%"
    sm.cell(t + 2, 1, "Solved rate = worked / (worked + didn't work). Pending answers are left out until the asker reacts.")

    # ---- Decisions
    dc = wb.create_sheet("Decisions")
    dh = [("id", "Question ID"), ("close_month", "Close month"), ("area", "Area"), ("question", "Question"),
          ("owner_answer", "Owner's answer"), ("answered_by", "Answered by"), ("decision", "Decision / what changes"),
          ("change_type", "Change type"), ("worked", "Worked?")]
    for j, (k, h) in enumerate(dh, 1):
        c = dc.cell(1, j, h)
        c.fill, c.font, c.alignment = HEAD, WHITE, WRAP
        dc.column_dimensions[get_column_letter(j)].width = WIDTH.get(k, 15)
    for i, r in enumerate([r for r in rows if r["owner_answer"]], 2):
        for j, (k, _) in enumerate(dh, 1):
            dc.cell(i, j, r[k] or None).alignment = WRAP
    dc.freeze_panes = "B2"

    # ---- Trail
    tr = wb.create_sheet("Trail")
    th = [("at", "When (UTC)", 17), ("question_id", "Question", 10), ("actor", "Who", 26),
          ("event", "What happened", 26), ("detail", "Detail", 110)]
    for j, (_, h, w) in enumerate(th, 1):
        c = tr.cell(1, j, h)
        c.fill, c.font = HEAD, WHITE
        tr.column_dimensions[get_column_letter(j)].width = w
    for i, e in enumerate(log.events(), 2):
        for j, (k, _, _) in enumerate(th, 1):
            tr.cell(i, j, e[k] or None).alignment = WRAP
    tr.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
