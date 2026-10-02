"""
What the agent can do. The same tools serve the Slack app, the MCP server and the simulator.

Read-only on the close:  close_overview, list_tables, run_sql, trace_cell, read_assumption
People:                  find_owner
Question log:            search_log, log_question, record_answer, escalate, record_owner_answer,
                         record_outcome

Nothing here can edit source data, config.yaml, the workbook or a check result. The warehouse is
opened read-only, file access from SQL is switched off, and queries naming request-content
columns are refused.
"""

import json
import re
from pathlib import Path

import duckdb
import yaml

from .qlog import CHANGE_TYPES, ESCALATION_REASONS, QuestionLog

ROOT = Path(__file__).resolve().parent.parent
MAX_ROWS = 200

# Which query and SQL file fill each Data_ sheet, so a traced cell ends at its source.
DATA_SOURCES = {
    "Data_Params": ("p, prices, source_status", "lib/run.py (params)"),
    "Data_Months": ("months, month_days", "sql/40_analytics.sql"),
    "Data_Monthly": ("monthly_summary", "sql/30_core.sql"),
    "Data_Lists": ("lists", "sql/40_analytics.sql"),
    "Data_UseCase": ("use_case_month", "sql/40_analytics.sql"),
    "Data_InferenceAcc": ("inference_accuracy", "sql/40_analytics.sql"),
    "Data_Cohort": ("cohort_tokens", "sql/40_analytics.sql"),
    "Data_EarlierCohorts": ("earlier_cohorts", "sql/40_analytics.sql"),
    "Data_Signups": ("cohort_signups", "sql/40_analytics.sql"),
    "Data_Contracts": ("contracts_live", "sql/40_analytics.sql"),
    "Data_EntMonth": ("enterprise_month", "sql/40_analytics.sql"),
    "Data_Pipeline": ("pipeline_open", "sql/40_analytics.sql"),
    "Data_Concentration": ("concentration", "sql/40_analytics.sql"),
    "Data_Checks": ("checks", "sql/50_checks.sql"),
    "Data_Reliability": ("reliability", "sql/40_analytics.sql"),
    "Data_Hourly": ("hourly_last7", "sql/40_analytics.sql"),
    "Data_FcHistory": ("forecast_history.csv", "lib/run.py + lib/forecast.py"),
}

COLREF = re.compile(r"(Data_[A-Za-z]+)!\$?([A-Z]{1,3}):\$?[A-Z]{1,3}(?![0-9$])")
REF = re.compile(r"(?<![A-Za-z0-9_.$'])(?:'?([A-Za-z_][A-Za-z0-9_]*)'?!)?\$?([A-Z]{1,3})\$?([0-9]+)(?::\$?([A-Z]{1,3})\$?([0-9]+))?")


def mention(person):
    """Slack mention text for {name, slack_id}. U/W = a person, S = a user group."""
    sid = (person or {}).get("slack_id") or ""
    name = (person or {}).get("name") or "someone"
    if sid.startswith(("U", "W")):
        return f"<@{sid}>"
    if sid.startswith("S"):
        return f"<!subteam^{sid}>"
    return f"@{name} (no Slack ID set)"


class Agent:
    """Holds the settings, the close outputs and the question log for one deployment."""

    def __init__(self, cfg_path=None, cfg=None):
        if cfg is None:
            p = Path(cfg_path or ROOT / "agent" / "agent.yaml")
            if not p.exists():
                p = ROOT / "agent" / "agent.example.yaml"
            cfg = yaml.safe_load(p.read_text())
        self.cfg = cfg
        out = Path(cfg.get("outputs_dir", "outputs"))
        self.out_dir = out if out.is_absolute() else ROOT / out
        lp = Path(cfg.get("log_path", "outputs/agent/questions.sqlite"))
        self.log = QuestionLog(lp if lp.is_absolute() else ROOT / lp)
        self.blocked = [w.lower() for w in cfg.get("blocked_columns", [])]

    # ------------------------------------------------------------ close files
    def close_month(self):
        m = str(self.cfg.get("close_month", "latest"))
        if m != "latest":
            return m
        files = sorted(self.out_dir.glob("close_*.json"))
        if not files:
            raise FileNotFoundError(f"No close found in {self.out_dir}. Run: python jevclose.py run")
        return files[-1].stem.replace("close_", "")

    def close_json(self):
        return json.loads((self.out_dir / f"close_{self.close_month()}.json").read_text())

    def workbook_path(self):
        return self.out_dir / f"Close_{self.close_month()}.xlsx"

    def _db(self):
        """Read-only warehouse. Source files behind the src_ views stay readable; nothing else on disk
        or the network is (DuckDB allowed_paths + enable_external_access = false)."""
        con = duckdb.connect(str(self.out_dir / "warehouse.duckdb"), read_only=True)
        paths, dirs = set(), set()
        for (sql,) in con.execute("SELECT sql FROM duckdb_views() WHERE NOT internal AND view_name LIKE 'raw_%'").fetchall():
            for m in re.finditer(r"read_\w+\((\[[^\]]*\]|'[^']*')", sql):
                for p in re.findall(r"'([^']+)'", m.group(1)):
                    if any(ch in p for ch in "*?["):
                        parts = Path(p).parts
                        stem = next(i for i, x in enumerate(parts) if any(ch in x for ch in "*?["))
                        dirs.add(str(Path(*parts[:stem])) + "/")
                    else:
                        paths.add(p)
        q = lambda xs: "[" + ", ".join("'" + x.replace("'", "''") + "'" for x in sorted(xs)) + "]"  # noqa: E731
        if paths:
            con.execute(f"SET allowed_paths = {q(paths)}")
        if dirs:
            con.execute(f"SET allowed_directories = {q(dirs)}")
        con.execute("SET enable_external_access = false")
        con.execute("SET lock_configuration = true")
        return con

    # ------------------------------------------------------------ owners
    def areas(self):
        return self.cfg.get("areas") or {}

    def owner_of(self, area):
        a = self.areas().get(area)
        if not a:
            return None, None
        if (a.get("owner") or {}).get("slack_id"):
            return a["owner"], "owner"
        if (a.get("backup") or {}).get("slack_id"):
            return a["backup"], "backup"
        return self.cfg.get("close_lead") or {"name": "Close lead"}, "close lead (no owner set)"


# ====================================================================== tools
def close_overview(ag, ctx):
    d = ag.close_json()
    checks = d["checks"]
    logged = {}
    for r in ag.log.all():
        if r["close_month"] == ag.close_month():
            logged.setdefault(r["question"].lower(), r["id"])
    attention = [c for c in checks if c["status"] in ("FLAG", "INFO")]
    counts = {}
    for c in checks:
        counts[c["status"]] = counts.get(c["status"], 0) + 1
    con = ag._db()
    rev = con.execute("""SELECT strftime(month, '%Y-%m') AS month, ROUND(SUM(billable_tokens)/1e9, 1) AS billable_btok,
                         ROUND(SUM(net_revenue), 0) AS net_revenue_usd FROM monthly_summary
                         GROUP BY 1 ORDER BY 1 DESC LIMIT 3""").fetchall()
    con.close()
    fc = {k: round(v, 1) for k, v in d.get("forecast_btok_per_day", {}).items() if k.startswith("base|")}
    return {
        "close_month": ag.close_month(), "usage_grain": d.get("grain"), "sources": d.get("sources"),
        "workbook": Path(d.get("workbook", "")).name,
        "check_counts": counts,
        "needs_attention": [{k: c[k] for k in ("check_no", "check_name", "scope", "status", "detail", "usd_impact")}
                            for c in attention],
        "last_3_months": [{"month": m, "billable_B_tokens": t, "net_revenue_usd": r} for m, t, r in rev],
        "forecast_base_B_tokens_per_day": fc,
        "questions_waiting_on_owners": [{k: r[k] for k in ("id", "question", "escalated_to", "opened_at")}
                                        for r in ag.log.waiting()],
    }


def list_tables(ag, ctx, names=None):
    con = ag._db()
    tables = [r[0] for r in con.execute("SHOW TABLES").fetchall() if not r[0].startswith("raw_")]
    if names:
        tables = [t for t in tables if t in names]
    out = {t: [f"{c} {ty}" for c, ty, *_ in con.execute(f"DESCRIBE {t}").fetchall()] for t in tables}
    con.close()
    return out


def run_sql(ag, ctx, sql):
    s = sql.strip().rstrip(";")
    if not re.match(r"(?is)^\s*(select|with|describe|summarize|show)\b", s) or ";" in s:
        return {"error": "Only a single SELECT / WITH / DESCRIBE query is allowed."}
    low = s.lower()
    hit = [w for w in ag.blocked if re.search(rf"\b{re.escape(w)}\b", low)]
    if hit:
        return {"error": f"Refused: queries may not touch request-content columns ({', '.join(hit)})."}
    if re.search(r"\b(read_\w+|glob|read_text|read_blob)\s*\(", low) or re.search(r"\bfrom\s+'", low):
        return {"error": "Query the warehouse tables, not files directly."}
    if re.search(r"\braw_", low):
        return {"error": "Query the src_ and model tables, not raw_ tables."}
    con = ag._db()
    try:
        cur = con.execute(s)
        cols = [c[0] for c in cur.description]
        rows = cur.fetchmany(MAX_ROWS + 1)
    except Exception as e:  # noqa: BLE001
        return {"error": str(e).split("\n")[0]}
    finally:
        con.close()
    trunc = len(rows) > MAX_ROWS
    rows = [[(round(v, 6) if isinstance(v, float) else (str(v) if v is not None and not isinstance(v, (int, str)) else v))
             for v in r] for r in rows[:MAX_ROWS]]
    return {"columns": cols, "rows": rows, "truncated": trunc, "sql": s}


def trace_cell(ag, ctx, sheet=None, cell=None, label=None):
    """Follow a workbook cell's formula back to the Data_ sheet and SQL that feed it."""
    from openpyxl import load_workbook
    wb = load_workbook(ag.workbook_path())
    if label and not (sheet and cell):
        hits = []
        for ws in wb.worksheets:
            if ws.title.startswith("Data_"):
                continue
            for row in ws.iter_rows(max_col=1):
                v = row[0].value
                if isinstance(v, str) and not v.startswith("=") and label.lower() in v.lower():
                    hits.append({"sheet": ws.title, "cell": f"B{row[0].row}", "label": v.strip()})
        return {"matches": hits[:15], "hint": "Call again with sheet and cell to trace one."}
    steps, seen = [], set()

    def walk(sh, ref, depth):
        key = (sh, ref)
        if key in seen or depth > 4 or len(steps) > 20:
            return
        seen.add(key)
        if sh not in wb.sheetnames:
            return
        ws = wb[sh]
        if sh.startswith("Data_"):
            col = re.match(r"[A-Z]+", ref).group()
            hdr_row = 1
            header = ws[f"{col}{hdr_row}"].value
            table, sql = DATA_SOURCES.get(sh, ("?", "?"))
            steps.append({"sheet": sh, "cell": ref, "column": header, "filled_from_table": table, "sql_file": sql})
            return
        v = ws[ref].value
        r = int(re.search(r"\d+", ref).group())
        lab = ws[f"A{r}"].value
        steps.append({"sheet": sh, "cell": ref, "row_label": lab if not str(lab).startswith("=") else None,
                      "formula": v if isinstance(v, str) and v.startswith("=") else None,
                      "value": None if isinstance(v, str) and v.startswith("=") else v})
        if isinstance(v, str) and v.startswith("="):
            refs = REF.findall(v)
            for dsh, dcol in dict.fromkeys(COLREF.findall(v)):
                walk(dsh, f"{dcol}1", depth + 1)
            cells = []
            for s2, c1, r1, c2, r2 in refs[:4]:
                if c2 and c1 == c2 and not (s2 or sh).startswith("Data_"):
                    cells += [(s2 or sh, f"{c1}{r}") for r in range(int(r1), min(int(r2), int(r1) + 3) + 1)]
                else:
                    cells.append((s2 or sh, f"{c1}{r1}"))
            for s2, ref2 in cells[:6]:
                walk(s2, ref2, depth + 1)

    walk(sheet, cell.replace("$", "").upper(), 0)
    return {"trace": steps,
            "note": "Formulas recalculate when the file opens. To get values, query the filled_from_table in run_sql."}


def read_assumption(ag, ctx, number):
    text = (ROOT / "LOGIC_REVIEW.md").read_text()
    m = re.search(rf"(?ms)^{int(number)}\. (.*?)(?=^\d+\. |^\*\*|^## )", text)
    return {"assumption": int(number), "text": m.group(1).strip() if m else "not found"}


def find_owner(ag, ctx, area=None, check_no=None, assumption_no=None, text=None):
    areas = ag.areas()
    pick = None
    if area and area in areas:
        pick = area
    elif check_no is not None:
        pick = next((k for k, a in areas.items() if int(check_no) in (a.get("checks") or [])), None)
    elif assumption_no is not None:
        pick = next((k for k, a in areas.items() if int(assumption_no) in (a.get("assumptions") or [])), None)
    elif text:
        t = text.lower()
        score = {k: sum(t.count(w.lower()) for w in (a.get("keywords") or [])) for k, a in areas.items()}
        best = max(score, key=score.get) if score else None
        pick = best if best and score[best] else None
    if not pick:
        return {"area": None, "note": "No area matched. Areas: " + ", ".join(areas), "close_lead": mention(ag.cfg.get("close_lead"))}
    person, role = ag.owner_of(pick)
    return {"area": pick, "area_name": areas[pick].get("name"), "who": person.get("name"),
            "mention": mention(person), "role": role}


def search_log(ag, ctx, text="", area=""):
    return {"matches": ag.log.search(text, area)}


def log_question(ag, ctx, question, area="", problem="", related_to=""):
    if area and area not in ag.areas():
        return {"error": f"area must be one of {list(ag.areas())}"}
    qid = ag.log.open(question=question, asked_by=ctx.get("asker", ""), close_month=ag.close_month(),
                      channel=ctx.get("channel", ""), thread_link=ctx.get("thread_link", ""), area=area,
                      problem=problem, related_to=related_to)
    return {"question_id": qid}


def record_answer(ag, ctx, question_id, answer, evidence=""):
    ag.log.update(question_id, "agent", "agent answered", agent_answer=answer, evidence=evidence, status="answered")
    ctx.setdefault("ask_feedback", []).append(question_id)
    return {"ok": True, "next": "Post the answer. The app will ask the asker whether it solved the problem."}


def escalate(ag, ctx, question_id, area, reason, question_for_owner, agent_answer_so_far=""):
    if reason not in ESCALATION_REASONS:
        return {"error": f"reason must be one of {ESCALATION_REASONS}"}
    o = find_owner(ag, ctx, area=area)
    if not o.get("area"):
        return {"error": f"unknown area {area}"}
    q = ag.log.get(question_id)
    fields = dict(status="waiting_on_owner", area=(q or {}).get("area") or area, escalation_reason=reason,
                  escalation_question=question_for_owner, escalated_to=f"{o['who']} {o['mention']}")
    if agent_answer_so_far:
        fields["agent_answer"] = agent_answer_so_far
    if q and q["escalated_to"]:
        prior = q["escalated_to"].split("; ")
        fields["escalated_to"] = "; ".join(dict.fromkeys(prior + [fields["escalated_to"]]))
    ag.log.update(question_id, "agent", f"asked {o['who']}", **fields)
    return {"post_this_mention": o["mention"], "owner": o["who"], "role": o["role"],
            "next": "Post one message in the thread that tags the owner and asks the question in one or two lines."}


def record_owner_answer(ag, ctx, question_id, answer, answered_by, decision="", change_type="none"):
    if change_type not in CHANGE_TYPES:
        return {"error": f"change_type must be one of {CHANGE_TYPES}"}
    from .qlog import now
    ag.log.update(question_id, answered_by, "owner answered", owner_answer=answer, answered_by=answered_by,
                  answered_at=now(), decision=decision, change_type=change_type, status="owner_answered")
    ctx.setdefault("ask_feedback", []).append(question_id)
    out = {"ok": True}
    if change_type in ("config_change", "input_cell_change", "source_data_fix", "billing_correction"):
        out["reminder"] = (f"This changes numbers. Do not make the change yourself. Tag the close lead "
                           f"{mention(ag.cfg.get('close_lead'))} to approve it, and say exactly what would change.")
    return out


def record_outcome(ag, ctx, question_id, worked, note=""):
    from .qlog import now
    fields = dict(worked=worked, worked_note=note)
    if worked == "yes":
        fields.update(status="closed", closed_at=now())
    ag.log.update(question_id, ctx.get("asker", ""), f"outcome: {worked}", **fields)
    return {"ok": True}


# ====================================================================== tool specs (JSON schema)
def _p(**props):
    req = [k for k, v in props.items() if v.pop("required", False)]
    return {"type": "object", "properties": props, "required": req}


S = lambda d, **kw: {"type": "string", "description": d, **kw}  # noqa: E731
I = lambda d, **kw: {"type": "integer", "description": d, **kw}  # noqa: E731

TOOLS = [
    ("close_overview", close_overview, "Start here. The close month, check results needing attention, last three months of tokens and revenue, next-month forecast, and questions still waiting on owners.", _p()),
    ("list_tables", list_tables, "Tables and columns you can query in the close warehouse.", _p(names={"type": "array", "items": {"type": "string"}, "description": "Only these tables (optional)"})),
    ("run_sql", run_sql, f"Run one read-only DuckDB SELECT on the close warehouse. Returns up to {MAX_ROWS} rows. Use the src_ tables for source data and the model tables (usage_daily, revenue_expected, monthly_summary, checks, contract_months, ...) for the close.", _p(sql=S("A single SELECT or WITH query", required=True))),
    ("trace_cell", trace_cell, "Follow a workbook cell back through its formulas to the Data_ sheet, table and SQL file it comes from. Give sheet+cell, or a label to find the cell first (e.g. 'Net revenue').", _p(sheet=S("Sheet name, e.g. Summary"), cell=S("Cell, e.g. B8"), label=S("Row label to search for"))),
    ("read_assumption", read_assumption, "Read assumption N from LOGIC_REVIEW.md (the 14 things the team must confirm).", _p(number=I("1-14", required=True))),
    ("find_owner", find_owner, "Who owns a topic. Give an area, a check number, an assumption number, or free text.", _p(area=S("Area key"), check_no=I("Check number"), assumption_no=I("Assumption number"), text=S("Free text to match against area keywords"))),
    ("search_log", search_log, "Search past questions, answers and owner decisions. Always do this before answering or escalating: the answer may already be settled.", _p(text=S("Keywords", required=True), area=S("Area key (optional)"))),
    ("log_question", log_question, "Open a log entry for the question you were just asked. Do this once per new question, before answering.", _p(question=S("The question, in the asker's words", required=True), area=S("Area key"), problem=S("One line: the problem behind the question"), related_to=S("ID of an earlier related question"))),
    ("record_answer", record_answer, "Record the answer you are about to post and the evidence (the SQL you ran or the cells you traced).", _p(question_id=S("Q-0001", required=True), answer=S("The answer, as posted", required=True), evidence=S("SQL, table or cell references"))),
    ("escalate", escalate, "Ask the owner of an area. Use when the answer is a judgment call, sources disagree, it contradicts a logged decision, a FLAG has no explanation, the answer would change numbers, or the data isn't there. Returns the mention to post.", _p(question_id=S("Q-0001", required=True), area=S("Area key", required=True), reason=S("Why", enum=list(ESCALATION_REASONS), required=True), question_for_owner=S("The exact question for the owner, one or two lines", required=True), agent_answer_so_far=S("What the data does show, if anything"))),
    ("record_owner_answer", record_owner_answer, "Record an owner's reply to an escalated question, the decision it implies, and whether it changes anything.", _p(question_id=S("Q-0001", required=True), answer=S("The owner's answer, summarized faithfully", required=True), answered_by=S("Owner's name or mention", required=True), decision=S("What this means for the close"), change_type=S("What changes", enum=list(CHANGE_TYPES)))),
    ("record_outcome", record_outcome, "Record whether the answer solved the problem, when the asker says so in words.", _p(question_id=S("Q-0001", required=True), worked=S("yes / no / pending", enum=["yes", "no", "pending"], required=True), note=S("What happened"))),
]

TOOL_FUNCS = {name: fn for name, fn, _, _ in TOOLS}


def anthropic_tools():
    return [{"name": n, "description": d, "input_schema": s} for n, _, d, s in TOOLS]


def call(ag, ctx, name, args):
    fn = TOOL_FUNCS.get(name)
    if not fn:
        return {"error": f"unknown tool {name}"}
    try:
        return fn(ag, ctx, **(args or {}))
    except (KeyError, ValueError, FileNotFoundError, TypeError) as e:
        return {"error": str(e)}
