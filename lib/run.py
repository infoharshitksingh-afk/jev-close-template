"""
The monthly close: sources -> DuckDB -> SQL -> Data_ tables -> workbook.

    python jevclose.py run [--as-of 2027-02] [--config config.yaml]
"""

import calendar
import csv
import json
from datetime import date, datetime
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl import load_workbook

from . import forecast as fc
from . import sources
from .schema import DATA_SHEETS, PARAMS, PARAM_ROW, WINDOW

ROOT = Path(__file__).resolve().parent.parent
SQL = ROOT / "sql"
TEMPLATE = ROOT / "template" / "Jev_Close_Template.xlsx"


# ------------------------------------------------------------------ helpers
def month_first(s):
    y, m = map(int, str(s)[:7].split("-"))
    return date(y, m, 1)


def month_end(d):
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def add_months(d, n):
    return fc.add_months(d, n)


def _num(x, default):
    return default if x is None else x


def params_sql(con, cfg, as_of):
    b = cfg.get("billing", {})
    ck = cfg.get("checks", {})
    uc = cfg.get("use_case", {})
    tz = cfg.get("reporting_timezone", "UTC")
    con.execute(f"""
        CREATE OR REPLACE TABLE p AS SELECT
            DATE '{as_of}'                                    AS as_of_month,
            DATE '{month_end(as_of)}'                         AS as_of_end,
            DATE '{add_months(as_of, -(WINDOW - 1))}'         AS window_start,
            '{tz}'                                            AS tz,
            {float(_num(b.get('free_credit_usd'), 0))}        AS free_credit_usd,
            {str(bool(b.get('prorate_partial_months', True))).upper()} AS prorate_partial_months,
            '{cfg.get('gateway_revenue_treatment', 'net')}'   AS gateway_revenue_treatment,
            {float(_num(ck.get('billed_tokens_tolerance_pct'), 0.0005))} AS tol_billed_pct,
            {float(_num(ck.get('gateway_tolerance_pct'), 0.005))}        AS tol_gateway_pct,
            {float(_num(ck.get('remittance_tolerance_usd'), 0.02))}      AS tol_remit_usd,
            {float(_num(ck.get('price_tolerance_per_mtok'), 1e-9))}      AS tol_price,
            {float(_num(ck.get('concentration_threshold'), 0.5))}        AS concentration_threshold,
            {float(_num(ck.get('unclassified_threshold'), 0.2))}         AS unclassified_threshold,
            {int(_num(uc.get('offhours_end_hour'), 6))}                  AS offhours_end_hour
    """)
    prices = b.get("prices") or [{"effective_from": "2000-01-01", "price_per_mtok": 0.042}]
    con.execute("CREATE OR REPLACE TABLE prices (effective_from DATE, price_per_mtok DOUBLE)")
    for pr in prices:
        con.execute("INSERT INTO prices VALUES (?, ?)", [str(pr["effective_from"]), float(pr["price_per_mtok"])])
    con.execute("CREATE OR REPLACE TABLE gateways (channel VARCHAR, fee_pct DOUBLE, aggregate BOOLEAN)")
    for ch, g in (cfg.get("gateways") or {}).items():
        con.execute("INSERT INTO gateways VALUES (?, ?, ?)",
                    [ch.lower(), float(g.get("fee_pct", 0)), not bool(g.get("passes_customer_ids", True))])
    con.execute("CREATE OR REPLACE TABLE billable_status (status INTEGER)")
    for st in b.get("billable_statuses", [200]):
        con.execute("INSERT INTO billable_status VALUES (?)", [int(st)])
    con.execute("CREATE OR REPLACE TABLE free_credit_channels (channel VARCHAR)")
    for ch in b.get("free_credit_channels", ["direct"]):
        con.execute("INSERT INTO free_credit_channels VALUES (?)", [ch.lower()])


def run_sql(con, grain):
    files = [SQL / f"10_usage_{grain}.sql", SQL / "20_use_case_rules.sql", SQL / "30_core.sql",
             SQL / "40_analytics.sql", SQL / "50_checks.sql"]
    for f in files:
        con.execute(f.read_text())


def detect_as_of(con, grain, tz):
    if grain == "call":
        mx = con.execute(f"SELECT MAX(CAST(timezone('{tz}', ts) AS DATE)) FROM src_usage").fetchone()[0]
    else:
        mx = con.execute("SELECT MAX(date) FROM src_usage").fetchone()[0]
    if mx is None:
        raise ValueError("The usage source has no rows with a valid date.")
    return month_first(mx) if mx == month_end(mx) else add_months(month_first(mx), -1)


def export(con, cfg, grain, status, as_of, include_pipeline=True):
    q = lambda s: con.execute(s).df()  # noqa: E731
    d = {}
    total_rev = con.execute("SELECT COALESCE(SUM(net_revenue), 0) FROM concentration").fetchone()[0]
    first, last = con.execute("SELECT MIN(date), MAX(date) FROM usage_events").fetchone()
    agg_fee = con.execute("SELECT COALESCE(AVG(fee_pct), 0) FROM gateways WHERE aggregate").fetchone()[0]
    price_now = con.execute(f"SELECT price_per_mtok FROM prices WHERE effective_from <= DATE '{month_end(as_of)}' "
                            "ORDER BY effective_from DESC LIMIT 1").fetchone()
    vals = {
        "as_of_month": as_of, "window_start": add_months(as_of, -(WINDOW - 1)),
        "list_price_per_mtok": price_now[0] if price_now else None,
        "free_credit_usd": float((cfg.get("billing") or {}).get("free_credit_usd", 0)),
        "aggregate_gateway_fee": agg_fee, "gateway_revenue_treatment": cfg.get("gateway_revenue_treatment", "net"),
        "total_net_revenue_asof": total_rev, "usage_grain": grain, "data_first_date": first, "data_last_date": last,
        "sources_loaded": ", ".join(k for k, v in status.items() if v == "loaded"),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "company_name": cfg.get("company_name", "Company"), "reporting_timezone": cfg.get("reporting_timezone", "UTC"),
    }
    d["Data_Params"] = pd.DataFrame([{"key": k, "value": vals[k], "description": desc} for k, desc in PARAMS])
    d["Data_Months"] = q("SELECT month, days, calendar_days FROM month_days ORDER BY month")
    d["Data_Monthly"] = q("""SELECT m.month, m.segment, m.channel, m.active_accounts, m.paying_accounts, m.billable_tokens,
                                    m.free_credit_tokens, m.free_credit_usd, m.list_value, m.usage_revenue, m.true_up,
                                    m.net_revenue, m.gateway_fees
                             FROM monthly_summary m, p WHERE m.month >= p.window_start ORDER BY 1, 2, 3""")
    d["Data_Lists"] = q("SELECT n, channel, use_case FROM lists ORDER BY n")
    d["Data_UseCase"] = q("SELECT month, use_case, billable_tokens, tokens_from_label FROM use_case_month ORDER BY 1, 2")
    d["Data_InferenceAcc"] = q("SELECT use_case, labelled_tokens, inferred_correctly FROM inference_accuracy ORDER BY labelled_tokens DESC")
    d["Data_Cohort"] = q("SELECT cohort, k, billable_tokens, active_accounts FROM cohort_tokens ORDER BY 1, 2")
    d["Data_EarlierCohorts"] = q("SELECT month, billable_tokens FROM earlier_cohorts ORDER BY 1")
    d["Data_Signups"] = q("SELECT cohort, signups, added_payment_method, payment_flag_known FROM cohort_signups ORDER BY 1")
    d["Data_Contracts"] = q("""SELECT customer, account_id, start_date, end_date, monthly_commit_usd, unit_price_per_mtok,
                                      discount_pct, last7_btok_per_day FROM contracts_live""")
    d["Data_EntMonth"] = q("SELECT customer, month, billable_tokens, usage_revenue, commitment, true_up FROM enterprise_month ORDER BY 1, 2")
    d["Data_Pipeline"] = (q("""SELECT deal_id, customer, stage, probability, expected_monthly_commit_usd, expected_start,
                                      expected_discount_pct, use_cases FROM pipeline_open""")
                          if include_pipeline else pd.DataFrame(columns=DATA_SHEETS["Data_Pipeline"]))
    d["Data_Concentration"] = q("""SELECT rank, account_id, display_name, segment, channel, billable_tokens, net_revenue,
                                          share_of_revenue FROM concentration ORDER BY rank""")
    d["Data_Checks"] = q("""SELECT check_no, check_name, scope, expected, actual, usd_impact, status, detail FROM checks
                            ORDER BY CASE status WHEN 'FLAG' THEN 0 WHEN 'INFO' THEN 1 WHEN 'SKIP' THEN 2 ELSE 3 END, check_no, scope""")
    d["Data_Reliability"] = q("SELECT month, calls_ok, calls_failed FROM reliability ORDER BY 1")
    d["Data_Hourly"] = q("SELECT hour_start, tokens FROM hourly_last7 ORDER BY 1")
    for name, cols in DATA_SHEETS.items():
        if name in d and name != "Data_FcHistory":
            d[name] = d[name][cols] if len(d[name]) else pd.DataFrame(columns=cols)
    return d


def forecast_inputs(cfg):
    f = {"tail_factor": {"low": 0.93, "base": 0.97, "high": 1.0}, "gateway_growth_decay": 0.5,
         "ent_usage_ratio": {"low": 0.8, "base": 1.0, "high": 1.2}, "ent_growth": {"low": 0.0, "base": 0.03, "high": 0.06},
         "ent_ramp_days": 60}
    for k, v in (cfg.get("forecast") or {}).items():
        if isinstance(v, dict) and isinstance(f.get(k), dict):
            f[k].update(v)
        else:
            f[k] = v
    return f


# ------------------------------------------------------------------ forecast history
HIST_COLS = ["close_month", "target_month", "scenario", "component", "btok_per_day", "source", "run_at"]


def load_history(path):
    if not path.exists():
        return []
    with open(path) as f:
        return list(csv.DictReader(f))


def save_history(path, rows, close, results, source):
    keep = [r for r in rows if r["close_month"] != str(close)]
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for (s, tm), comp in results.items():
        for c in ("self_serve", "gateway_aggregate", "enterprise", "total"):
            keep.append(dict(close_month=str(close), target_month=str(tm), scenario=s, component=c,
                             btok_per_day=round(comp[c], 6), source=source, run_at=now))
    keep.sort(key=lambda r: (r["close_month"], r["target_month"], r["scenario"], r["component"]))
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=HIST_COLS)
        w.writeheader()
        w.writerows(keep)
    return keep


# ------------------------------------------------------------------ workbook
def _cell_value(v):
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime() if (v.hour or v.minute) else v.date()
    if hasattr(v, "item"):
        return v.item()
    return v


def fill_workbook(data, cfg, inputs, out_path, input_cells):
    wb = load_workbook(TEMPLATE)
    for name, cols in DATA_SHEETS.items():
        ws = wb[name]
        if name == "Data_Params":
            for k, _ in PARAMS:
                v = data[name].loc[data[name]["key"] == k, "value"].iloc[0]
                c = ws.cell(row=PARAM_ROW[k], column=2, value=_cell_value(v))
            continue
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row)
        df = data[name]
        for i, row in enumerate(df.itertuples(index=False), start=2):
            for j, v in enumerate(row, start=1):
                c = ws.cell(row=i, column=j, value=_cell_value(v))
                if isinstance(c.value, (date, datetime)):
                    c.number_format = "yyyy-mm-dd hh:mm" if isinstance(c.value, datetime) else "yyyy-mm-dd"
    # seed judgment inputs from config
    cap = cfg.get("capacity") or {}
    for key, (sheet, ref) in input_cells.items():
        if key.startswith("capacity."):
            v = cap.get(key.split(".", 1)[1])
        elif "." in key:
            a, b = key.split(".")
            v = (inputs.get(a) or {}).get(b)
        else:
            v = inputs.get(key)
        if v is not None:
            wb[sheet][ref].value = v
    wb.calculation.fullCalcOnLoad = True  # Excel recalculates every formula when the file opens
    wb.save(out_path)


# ------------------------------------------------------------------ main entry
def close(cfg, base_dir, as_of=None, out_dir=None, verbose=True):
    say = print if verbose else (lambda *a, **k: None)
    out_dir = Path(out_dir or cfg.get("outputs", {}).get("dir", "outputs"))
    if not out_dir.is_absolute():
        out_dir = Path(base_dir) / out_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    db_path = out_dir / "warehouse.duckdb"
    if db_path.exists():
        db_path.unlink()
    con = duckdb.connect(str(db_path))
    status, grain = sources.load_all(con, cfg, base_dir)
    con.execute("CREATE OR REPLACE TABLE source_status AS SELECT * FROM (VALUES "
                + ", ".join(f"('{k}', {str(v == 'loaded').upper()})" for k, v in status.items()) + ") t(source, loaded)")
    say(f"Sources: {', '.join(f'{k}={v}' for k, v in status.items())}. Usage grain: {grain}.")
    tz = cfg.get("reporting_timezone", "UTC")
    as_of = month_first(as_of) if as_of else detect_as_of(con, grain, tz)
    inputs = forecast_inputs(cfg)
    hist_path = out_dir / "forecast_history.csv"
    hist = load_history(hist_path)

    prev = add_months(as_of, -1)
    if not any(r["close_month"] == str(prev) for r in hist):
        params_sql(con, cfg, prev)
        run_sql(con, grain)
        if con.execute("SELECT COUNT(*) FROM usage_daily").fetchone()[0] > 0:
            dprev = export(con, cfg, grain, status, prev, include_pipeline=False)
            hist = save_history(hist_path, hist, prev, fc.run(dprev, inputs, prev), "reconstructed (no pipeline)")
            say(f"Reconstructed the {prev:%b %Y} close's forecast for the variance check.")

    params_sql(con, cfg, as_of)
    run_sql(con, grain)
    data = export(con, cfg, grain, status, as_of)
    results = fc.run(data, inputs, as_of)
    hist = save_history(hist_path, hist, as_of, results, "live")
    data["Data_FcHistory"] = pd.DataFrame(
        [{"close_month": month_first(r["close_month"]), "target_month": month_first(r["target_month"]),
          "scenario": r["scenario"], "component": r["component"], "btok_per_day": float(r["btok_per_day"])} for r in hist],
        columns=DATA_SHEETS["Data_FcHistory"])

    if not TEMPLATE.exists() or not TEMPLATE.with_name("input_cells.json").exists():
        from .workbook import build
        build(TEMPLATE)
    input_cells = json.loads(TEMPLATE.with_name("input_cells.json").read_text())
    xlsx = out_dir / f"Close_{as_of:%Y-%m}.xlsx"
    fill_workbook(data, cfg, inputs, xlsx, input_cells)

    csv_dir = out_dir / f"data_{as_of:%Y-%m}"
    csv_dir.mkdir(exist_ok=True)
    for name, df in data.items():
        df.to_csv(csv_dir / f"{name}.csv", index=False)
    checks = data["Data_Checks"]
    summary = {
        "as_of": str(as_of), "grain": grain, "sources": status, "workbook": str(xlsx),
        "checks": json.loads(checks.to_json(orient="records", date_format="iso")),
        "forecast_btok_per_day": {f"{s}|{tm}": round(v["total"], 4) for (s, tm), v in results.items()},
        "forecast_revenue_k": {f"{s}|{tm}": round(v["revenue_k"], 2) for (s, tm), v in results.items()},
    }
    (out_dir / f"close_{as_of:%Y-%m}.json").write_text(json.dumps(summary, indent=1, default=str))
    con.close()
    n_flag = int((checks["status"] == "FLAG").sum()) if len(checks) else 0
    say(f"Close {as_of:%b %Y}: {n_flag} checks flagged. Workbook: {xlsx}")
    return xlsx, summary
