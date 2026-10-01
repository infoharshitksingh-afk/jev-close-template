"""
Build the formula-only workbook template.

Rule: numbers enter ONLY through the Data_ sheets (written by the runner, or pasted by hand).
Every cell on a report tab is a formula, a label, or a blue input. Report tabs read a fixed
12-month window that ends at the as-of month in Data_Params, so the same template works for
any month and any company size.

    python jevclose.py template      # rebuilds template/Jev_Close_Template.xlsx
"""

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as CL

from .schema import (DATA_SHEETS, MAX_CHANNELS, MAX_CHECK_ROWS, MAX_CONC, MAX_CONTRACTS, MAX_DEALS,
                     MAX_USE_CASES, PARAMS, P, WINDOW)

F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=14, bold=True)
SUB = Font(name=F, size=9, italic=True, color="555555")
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
RED_B = Font(name=F, size=11, bold=True, color="C00000")
HFILL = PatternFill("solid", fgColor="1F3A5F")
DFILL = PatternFill("solid", fgColor="5B6770")
SEC = PatternFill("solid", fgColor="E8EEF5")
YEL = PatternFill("solid", fgColor="FFFF00")
INFILL = PatternFill("solid", fgColor="F3F6FA")
WRAP = Alignment(wrap_text=True, vertical="top")

NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
NUM2 = '#,##0.00;(#,##0.00);-'
NUM3 = '#,##0.000;(#,##0.000);-'
NUM6 = '0.000000'
USD = '$#,##0;($#,##0);-'
USD2 = '$#,##0.00;($#,##0.00);-'
USD4 = '$#,##0.0000;($#,##0.0000);-'
PCT = '0.0%;(0.0%);-'
PCTR = '0%;[Red](0%);-'
MULT = '0.00"x";(0.00"x");-'
DATEF = 'yyyy-mm-dd'
MONF = 'mmm yyyy'

ASOF = P("as_of_month")
WSTART = P("window_start")

# Cells the runner seeds from config (the team can still edit them in Excel afterwards).
INPUT_CELLS = {}


def put(ws, ref, v, font=BLACK, fmt=None, fill=None, wrap=False):
    c = ws[ref]
    c.value = v
    c.font = font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    return c


def header(ws, row, labels, col=1, fill=HFILL):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font = HDR
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if isinstance(lab, str) and lab.startswith("="):
            c.number_format = MONF


def title(ws, t, s):
    put(ws, "A1", f'={P("company_name")}&" | {t}"', TITLE)
    put(ws, "A2", s, SUB)


def section(ws, row, text, ncols):
    for j in range(1, ncols + 1):
        ws.cell(row=row, column=j).fill = SEC
    ws.cell(row=row, column=1, value=text).font = BOLD


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def mon(j):
    """Window month j (0..11) as a formula."""
    return f"EDATE({WSTART},{j})"


def days_of(month_expr):
    return f"SUMIFS(Data_Months!$B:$B,Data_Months!$A:$A,{month_expr})"


def safe_days(month_expr):
    return f"MAX(1,{days_of(month_expr)})"


def ix(sheet, col, n):
    """Blank-safe INDEX into a Data_ sheet column (row n of the data, 1-based)."""
    ref = f"INDEX({sheet}!${col}:${col},{n}+1)"
    return f'IF({ref}="","",{ref})'


def build(path):
    wb = Workbook()

    # ======================================================================== Data sheets
    first = True
    for name, cols in DATA_SHEETS.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        header(ws, 1, cols, fill=DFILL)
        ws.sheet_properties.tabColor = "8A939B"
        ws.freeze_panes = "A2"
        for j, c in enumerate(cols):
            ws.column_dimensions[CL(1 + j)].width = max(12, len(c) + 4)
    dp = wb["Data_Params"]
    for i, (k, d) in enumerate(PARAMS):
        put(dp, f"A{i + 2}", k)
        put(dp, f"C{i + 2}", d, SUB)
    for k in ("as_of_month", "window_start", "data_first_date", "data_last_date"):
        dp[P(k).split("!")[1].replace("$", "")].number_format = DATEF
    widths(dp, {"A": 26, "B": 22, "C": 60})

    # ======================================================================== Monthly
    M = wb.create_sheet("Monthly", 0)
    title(M, "Usage and revenue by month", "Formulas only: every number reads the Data_ sheets. Tokens in billions, $ in thousands.")
    MC = [CL(2 + j) for j in range(WINDOW)]
    header(M, 4, ["Line"] + [f"={mon(j)}" for j in range(WINDOW)])
    segs = [("enterprise", "Enterprise"), ("self_serve", "Self-serve"),
            ("gateway_aggregate", "Gateways that don't pass customer IDs"), ("unmatched", "Unmatched accounts (fix the accounts source)")]

    def seg_sum(field_col, seg, j):
        return f"SUMIFS(Data_Monthly!${field_col}:${field_col},Data_Monthly!$A:$A,{MC[j]}$4,Data_Monthly!$B:$B,\"{seg}\")"

    def all_sum(field_col, j):
        return f"SUMIFS(Data_Monthly!${field_col}:${field_col},Data_Monthly!$A:$A,{MC[j]}$4)"

    r = 5
    section(M, r, "Billable tokens (B)", WINDOW + 1); r += 1
    tok_rows = []
    for seg, lab in segs:
        put(M, f"A{r}", lab)
        for j in range(WINDOW):
            put(M, f"{MC[j]}{r}", f"={seg_sum('F', seg, j)}/1E9", BLACK, NUM1)
        tok_rows.append(r); r += 1
    put(M, f"A{r}", "Total billable tokens", BOLD)
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=SUM({MC[j]}{tok_rows[0]}:{MC[j]}{tok_rows[-1]})", BOLD, NUM1)
    M_TOK = r; r += 1
    put(M, f"A{r}", "Days of data in month")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"={days_of(MC[j] + '$4')}", BLACK, NUM)
    M_DAYS = r; r += 1
    put(M, f"A{r}", "Tokens per day (B)", BOLD)
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=IF({MC[j]}{M_DAYS}=0,0,{MC[j]}{M_TOK}/{MC[j]}{M_DAYS})", BOLD, NUM1)
    M_TPD = r; r += 1
    seg_tpd_rows = {}
    for (seg, lab), tr in zip(segs[:3], tok_rows[:3]):
        put(M, f"A{r}", f"  {lab}, per day")
        for j in range(WINDOW):
            put(M, f"{MC[j]}{r}", f"=IF({MC[j]}{M_DAYS}=0,0,{MC[j]}{tr}/{MC[j]}{M_DAYS})", BLACK, NUM2)
        seg_tpd_rows[seg] = r; r += 1
    put(M, f"A{r}", "Growth in tokens per day vs prior month")
    for j in range(1, WINDOW):
        put(M, f"{MC[j]}{r}", f"=IF({MC[j - 1]}{M_TPD}=0,0,{MC[j]}{M_TPD}/{MC[j - 1]}{M_TPD}-1)", BLACK, PCT)
    M_GR = r; r += 2

    section(M, r, "Billable tokens by channel (B)", WINDOW + 1); r += 1
    for n in range(1, MAX_CHANNELS + 1):
        put(M, f"A{r}", f"={ix('Data_Lists', 'B', n)}")
        for j in range(WINDOW):
            put(M, f"{MC[j]}{r}", f'=IF($A{r}="","",SUMIFS(Data_Monthly!$F:$F,Data_Monthly!$A:$A,{MC[j]}$4,Data_Monthly!$C:$C,$A{r})/1E9)', BLACK, NUM1)
        r += 1
    r += 1

    section(M, r, "Net revenue ($k)", WINDOW + 1); r += 1
    rev_rows = []
    for seg, lab in segs:
        put(M, f"A{r}", lab)
        for j in range(WINDOW):
            put(M, f"{MC[j]}{r}", f"={seg_sum('L', seg, j)}/1E3", BLACK, NUM1)
        rev_rows.append(r); r += 1
    put(M, f"A{r}", "Total net revenue", BOLD)
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=SUM({MC[j]}{rev_rows[0]}:{MC[j]}{rev_rows[-1]})", BOLD, NUM1)
    M_REV = r; r += 1
    put(M, f"A{r}", "  of which enterprise true-ups")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"={all_sum('K', j)}/1E3", BLACK, NUM1)
    r += 1
    put(M, f"A{r}", "Annualized run-rate ($k)")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=IF({MC[j]}{M_DAYS}=0,0,{MC[j]}{M_REV}/{MC[j]}{M_DAYS}*365)", BLACK, NUM)
    M_ARR = r; r += 1
    put(M, f"A{r}", "Enterprise share of net revenue")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=IF({MC[j]}{M_REV}=0,0,{MC[j]}{rev_rows[0]}/{MC[j]}{M_REV})", BLACK, PCT)
    M_ENTSH = r; r += 1
    put(M, f"A{r}", "Realized net price per M tokens ($)")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"=IF({MC[j]}{M_TOK}=0,0,{MC[j]}{M_REV}/{MC[j]}{M_TOK})", BLACK, USD4)
    M_PRICE = r; r += 2

    section(M, r, "Other lines", WINDOW + 1); r += 1
    for lab, col, div, fmt in [("List value of tokens ($k)", "I", 1e3, NUM1), ("Free credit used ($k)", "H", 1e3, NUM2),
                               ("Gateway fees ($k)", "M", 1e3, NUM1)]:
        put(M, f"A{r}", lab)
        for j in range(WINDOW):
            put(M, f"{MC[j]}{r}", f"={all_sum(col, j)}/{div:.0f}", BLACK, fmt)
        r += 1
    put(M, f"A{r}", "Active accounts (gateways counted once)")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"={all_sum('D', j)}", BLACK, NUM)
    M_ACT = r; r += 1
    put(M, f"A{r}", "Paying accounts")
    for j in range(WINDOW):
        put(M, f"{MC[j]}{r}", f"={all_sum('E', j)}", BLACK, NUM)
    M_PAY = r; r += 1
    put(M, f"A{r}", "Failed calls as % of all calls (not billed)")
    for j in range(WINDOW):
        ok = f"SUMIFS(Data_Reliability!$B:$B,Data_Reliability!$A:$A,{MC[j]}$4)"
        bad = f"SUMIFS(Data_Reliability!$C:$C,Data_Reliability!$A:$A,{MC[j]}$4)"
        put(M, f"{MC[j]}{r}", f"=IF({ok}+{bad}=0,0,{bad}/({ok}+{bad}))", BLACK, PCT)
    M_FAIL = r
    widths(M, {"A": 46, **{c: 11 for c in MC}})
    M.freeze_panes = "B5"

    # ======================================================================== UseCases
    U = wb.create_sheet("UseCases", 1)
    title(U, "Where usage comes from", "Use case from the customer's project label when mapped (label_map), otherwise inferred from traffic shape (sql/20_use_case_rules.sql).")
    UC = [CL(2 + j) for j in range(WINDOW)]
    header(U, 4, ["Use case (B tokens)"] + [f"={mon(j)}" for j in range(WINDOW)] + ["Share of latest month", "From label, latest month"])
    u0 = 5
    for n in range(1, MAX_USE_CASES + 1):
        rr = u0 + n - 1
        put(U, f"A{rr}", f"={ix('Data_Lists', 'C', n)}")
        for j in range(WINDOW):
            put(U, f"{UC[j]}{rr}", f'=IF($A{rr}="","",SUMIFS(Data_UseCase!$C:$C,Data_UseCase!$A:$A,{UC[j]}$4,Data_UseCase!$B:$B,$A{rr})/1E9)', BLACK, NUM1)
        last = UC[-1]
        put(U, f"{CL(2 + WINDOW)}{rr}", f'=IF($A{rr}="","",IF({last}${u0 + MAX_USE_CASES}=0,0,{last}{rr}/{last}${u0 + MAX_USE_CASES}))', BLACK, PCT)
        lab = f"SUMIFS(Data_UseCase!$D:$D,Data_UseCase!$A:$A,{last}$4,Data_UseCase!$B:$B,$A{rr})/1E9"
        put(U, f"{CL(3 + WINDOW)}{rr}", f'=IF(OR($A{rr}="",N({last}{rr})=0),"",{lab}/{last}{rr})', BLACK, PCT)
    U_TOT = u0 + MAX_USE_CASES
    put(U, f"A{U_TOT}", "Total", BOLD)
    for j in range(WINDOW):
        put(U, f"{UC[j]}{U_TOT}", f"=SUM({UC[j]}{u0}:{UC[j]}{U_TOT - 1})", BOLD, NUM1)
    r = U_TOT + 2
    section(U, r, "How reliable is the inferred use case? Rules scored on traffic whose label we know", 6); r += 1
    header(U, r, ["Use case", "Labelled tokens (B)", "Inferred correctly (B)", "Accuracy"]); r += 1
    a0 = r
    for n in range(1, MAX_USE_CASES + 1):
        put(U, f"A{r}", f"={ix('Data_InferenceAcc', 'A', n)}")
        put(U, f"B{r}", f'=IF($A{r}="","",INDEX(Data_InferenceAcc!$B:$B,{n}+1)/1E9)', BLACK, NUM1)
        put(U, f"C{r}", f'=IF($A{r}="","",INDEX(Data_InferenceAcc!$C:$C,{n}+1)/1E9)', BLACK, NUM1)
        put(U, f"D{r}", f'=IF(OR($A{r}="",N(B{r})=0),"",C{r}/B{r})', BLACK, PCT)
        r += 1
    put(U, f"A{r}", "All labelled traffic", BOLD)
    put(U, f"B{r}", f"=SUM(B{a0}:B{r - 1})", BOLD, NUM1)
    put(U, f"C{r}", f"=SUM(C{a0}:C{r - 1})", BOLD, NUM1)
    put(U, f"D{r}", f'=IF(B{r}=0,"no labelled traffic",C{r}/B{r})', BOLD, PCT)

    put(U, f"A{r + 2}", "Low accuracy means the rules in sql/20_use_case_rules.sql don't fit your traffic yet. Recalibrate them, or map more labels.", SUB)
    widths(U, {"A": 30, **{c: 11 for c in UC}, CL(2 + WINDOW): 14, CL(3 + WINDOW): 14})

    # ======================================================================== Cohorts
    K = wb.create_sheet("Cohorts", 2)
    title(K, "Self-serve cohorts", "Accounts grouped by signup month. Tokens per day (B) so partial months compare fairly. Older cohorts are combined on the Forecast tab.")
    KC = [CL(4 + k) for k in range(WINDOW)]
    header(K, 4, ["Signup cohort", "Signups", "Added payment method"] + [f"Month {k}" for k in range(WINDOW)])
    for i in range(WINDOW):
        rr = 5 + i
        put(K, f"A{rr}", f"={mon(i)}", BLACK, MONF)
        put(K, f"B{rr}", f"=SUMIFS(Data_Signups!$B:$B,Data_Signups!$A:$A,$A{rr})", BLACK, NUM)
        put(K, f"C{rr}", f'=IF(SUMIFS(Data_Signups!$D:$D,Data_Signups!$A:$A,$A{rr})=0,"n/a",SUMIFS(Data_Signups!$C:$C,Data_Signups!$A:$A,$A{rr}))', BLACK, NUM)
        for k in range(WINDOW - i):
            tok = f"SUMIFS(Data_Cohort!$C:$C,Data_Cohort!$A:$A,$A{rr},Data_Cohort!$B:$B,{k})"
            put(K, f"{KC[k]}{rr}", f"={tok}/{safe_days(f'EDATE($A{rr},{k})')}/1E9", BLACK, NUM3)
    r = 5 + WINDOW + 1
    put(K, f"A{r}", "Active accounts as % of signups (retention)", BOLD); r += 1
    header(K, r, ["Signup cohort", "Signups", ""] + [f"Month {k}" for k in range(WINDOW)]); r += 1
    for i in range(WINDOW):
        put(K, f"A{r}", f"=A{5 + i}", BLACK, MONF)
        put(K, f"B{r}", f"=B{5 + i}", BLACK, NUM)
        for k in range(WINDOW - i):
            act = f"SUMIFS(Data_Cohort!$D:$D,Data_Cohort!$A:$A,$A{r},Data_Cohort!$B:$B,{k})"
            put(K, f"{KC[k]}{r}", f'=IF($B{r}=0,"",{act}/$B{r})', BLACK, PCT)
        r += 1
    put(K, f"A{r + 1}", "Read across a row to follow one cohort; read down a column to compare cohorts at the same age.", SUB)
    widths(K, {"A": 14, "B": 10, "C": 14, **{c: 10 for c in KC}})

    # ======================================================================== Enterprise
    E = wb.create_sheet("Enterprise", 3)
    title(E, "Enterprise: usage vs commitment", "Contracts live at the close (Data_Contracts) and usage revenue by month (Data_EntMonth). True-up = commitment minus usage when usage falls short.")
    EC = [CL(8 + j) for j in range(WINDOW)]
    header(E, 4, ["Customer", "Start", "End", "Monthly commit ($)", "Price ($/M)", "Latest month vs commit", "Latest month true-up ($)"]
           + [f"={mon(j)}" for j in range(WINDOW)])
    e0 = 5
    for n in range(1, MAX_CONTRACTS + 1):
        rr = e0 + n - 1
        put(E, f"A{rr}", f"={ix('Data_Contracts', 'A', n)}")
        put(E, f"B{rr}", f"={ix('Data_Contracts', 'C', n)}", BLACK, DATEF)
        put(E, f"C{rr}", f"={ix('Data_Contracts', 'D', n)}", BLACK, DATEF)
        put(E, f"D{rr}", f"={ix('Data_Contracts', 'E', n)}", BLACK, USD)
        put(E, f"E{rr}", f"={ix('Data_Contracts', 'F', n)}", BLACK, USD4)
        for j in range(WINDOW):
            put(E, f"{EC[j]}{rr}", f'=IF($A{rr}="","",SUMIFS(Data_EntMonth!$D:$D,Data_EntMonth!$A:$A,$A{rr},Data_EntMonth!$B:$B,{EC[j]}$4))', BLACK, USD)
        last = EC[-1]
        put(E, f"F{rr}", f'=IF(OR($A{rr}="",N($D{rr})=0),"",{last}{rr}/$D{rr})', BLACK, PCT)
        put(E, f"G{rr}", f'=IF($A{rr}="","",SUMIFS(Data_EntMonth!$F:$F,Data_EntMonth!$A:$A,$A{rr},Data_EntMonth!$B:$B,{last}$4))', BLACK, USD)
    E_TOT = e0 + MAX_CONTRACTS
    put(E, f"A{E_TOT}", "Total", BOLD)
    for col in ["D", "G"] + EC:
        put(E, f"{col}{E_TOT}", f"=SUM({col}{e0}:{col}{E_TOT - 1})", BOLD, USD)
    r = E_TOT + 3
    section(E, r, "Open pipeline (CRM)", 8); r += 1
    header(E, r, ["Deal", "Customer", "Stage", "Probability", "Expected monthly commit ($)", "Expected start", "Weighted commit ($)", "Use cases"]); r += 1
    p0 = r
    for n in range(1, MAX_DEALS + 1):
        for col, src, fmt in [("A", "A", None), ("B", "B", None), ("C", "C", None), ("D", "D", PCT), ("E", "E", USD), ("F", "F", DATEF), ("H", "H", None)]:
            put(E, f"{col}{r}", f"={ix('Data_Pipeline', src, n)}", BLACK, fmt)
        put(E, f"G{r}", f'=IF($A{r}="","",D{r}*E{r})', BLACK, USD)
        r += 1
    put(E, f"A{r}", "Total", BOLD)
    put(E, f"E{r}", f"=SUM(E{p0}:E{r - 1})", BOLD, USD)
    put(E, f"G{r}", f"=SUM(G{p0}:G{r - 1})", BOLD, USD)
    widths(E, {"A": 24, "B": 18, "C": 16, "D": 14, "E": 14, "F": 14, "G": 14, "H": 26, **{c: 11 for c in EC}})
    E.freeze_panes = "B5"

    # ======================================================================== Concentration
    N_ = wb.create_sheet("Concentration", 4)
    title(N_, "Customer concentration, latest month", "Top accounts by net revenue. A gateway that doesn't pass customer IDs counts as one account.")
    header(N_, 4, ["Rank", "Account", "Segment", "Channel", "Billable tokens (B)", "Net revenue ($)", "Share of revenue", "Cumulative share"])
    for n in range(1, MAX_CONC + 1):
        rr = 4 + n
        put(N_, f"A{rr}", f"={ix('Data_Concentration', 'A', n)}")
        put(N_, f"B{rr}", f"={ix('Data_Concentration', 'C', n)}")
        put(N_, f"C{rr}", f"={ix('Data_Concentration', 'D', n)}")
        put(N_, f"D{rr}", f"={ix('Data_Concentration', 'E', n)}")
        put(N_, f"E{rr}", f'=IF($A{rr}="","",INDEX(Data_Concentration!$F:$F,{n}+1)/1E9)', BLACK, NUM1)
        put(N_, f"F{rr}", f"={ix('Data_Concentration', 'G', n)}", BLACK, USD)
        put(N_, f"G{rr}", f'=IF(OR($A{rr}="",{P("total_net_revenue_asof")}=0),"",F{rr}/{P("total_net_revenue_asof")})', BLACK, PCT)
        put(N_, f"H{rr}", f'=IF($A{rr}="","",SUM($G$5:G{rr}))', BLACK, PCT)
    put(N_, "J4", "Net revenue, all accounts ($)", BOLD)
    put(N_, "J5", f"={P('total_net_revenue_asof')}", GREEN, USD)
    widths(N_, {"A": 6, "B": 36, "C": 18, "D": 14, "E": 16, "F": 16, "G": 14, "H": 14, "J": 30})

    # ======================================================================== Checks
    C = wb.create_sheet("Checks", 0)
    title(C, "Close checks", "Rules in sql/50_checks.sql (Data_Checks). FLAG = a person must look. Columns I-L are for people and the agent; they are not overwritten by formulas.")
    header(C, 4, ["#", "Check", "Scope", "Expected", "Actual", "$ impact", "Status", "Detail", "Owner",
                  "Next step", "Explanation (agent draft)", "Reviewed by"])
    for n in range(1, MAX_CHECK_ROWS + 1):
        rr = 4 + n
        for col, src, fmt in [("A", "A", None), ("B", "B", None), ("C", "C", None), ("D", "D", '#,##0.00;(#,##0.00);0'),
                              ("E", "E", '#,##0.00;(#,##0.00);0'), ("F", "F", USD2), ("G", "G", None), ("H", "H", None)]:
            put(C, f"{col}{rr}", f"={ix('Data_Checks', src, n)}", BOLD if col == "G" else BLACK, fmt)
        for col in "IJKL":
            C[f"{col}{rr}"].fill = INFILL
    last = 4 + MAX_CHECK_ROWS
    for val, color in (('"FLAG"', "F8D7DA"), ('"INFO"', "FFF3CD"), ('"PASS"', "D4EDDA"), ('"SKIP"', "E2E3E5")):
        C.conditional_formatting.add(f"G5:G{last}", CellIsRule(operator="equal", formula=[val], fill=PatternFill("solid", fgColor=color)))
    widths(C, {"A": 4, "B": 40, "C": 24, "D": 14, "E": 14, "F": 12, "G": 8, "H": 52, "I": 16, "J": 36, "K": 44, "L": 14})
    C.freeze_panes = "A5"

    # ======================================================================== Forecast
    FS = wb.create_sheet("Forecast")
    title(FS, "Demand forecast: next three months", "Self-serve by cohort (chain-ladder) + gateway trend + enterprise contracts and pipeline (Forecast_Ent). Blue cells are judgment inputs, seeded from config.")
    SC = {"low": "B", "base": "C", "high": "D"}
    SCN = ("low", "base", "high")
    put(FS, "A4", "Assumptions", BOLD)
    header(FS, 5, ["Input", "Low", "Base", "High", "Why"])
    r_tail, r_sg, r_m0, r_gw, r_decay, r_ratio, r_eg, r_ramp, r_ssp, r_gap, r_early = range(6, 17)
    labels = {
        r_tail: ("Cohort development factor beyond observed history", "Applies to cohort ages not yet observed, and to the combined older cohorts."),
        r_sg: ("New signups, monthly growth", "Base = average growth over the last 3 months. Low = flat (blue). High = double."),
        r_m0: ("Month-0 tokens/day per signup (B)", "Min / average / max of the last 3 cohorts."),
        r_gw: ("Gateway (no customer IDs) monthly growth", "From last month's growth: low = min(0, it), base = it x decay, high = max(0, it)."),
        r_decay: ("Gateway growth decay (base)", "Judgment: share of last month's growth expected to continue."),
        r_ratio: ("Enterprise steady-state usage vs commitment", "For contracts still ramping. Replace with the account team's view."),
        r_eg: ("Enterprise monthly growth after ramp", ""),
        r_ramp: ("Enterprise ramp length (days)", "Contracts younger than this ramp toward steady state."),
        r_ssp: ("Self-serve realized $ per M tokens (latest month)", "Net revenue / billable tokens. Below list price because of free credit and gateway fees."),
        r_gap: ("Gateway (no IDs) realized $ per M tokens (latest month)", ""),
        r_early: ("Older cohorts (before window), tokens/day latest month (B)", "Forecast = this x development factor each month."),
    }
    for rr, (lab, why) in labels.items():
        put(FS, f"A{rr}", lab)
        put(FS, f"E{rr}", why, SUB)

    W = WINDOW; T = W - 1
    TC = lambda k: CL(4 + k)  # noqa: E731
    A0 = 20
    put(FS, f"A{A0 - 1}", "Self-serve cohort triangle, actuals (B tokens per day)", BOLD)
    header(FS, A0, ["Cohort", "Signups", "Month-0 per signup"] + [f"Month {k}" for k in range(W)])
    for i in range(W):
        rr = A0 + 1 + i
        put(FS, f"A{rr}", f"={mon(i)}", BLACK, MONF)
        put(FS, f"B{rr}", f"=SUMIFS(Data_Signups!$B:$B,Data_Signups!$A:$A,$A{rr})", BLACK, NUM)
        put(FS, f"C{rr}", f'=IF(B{rr}>0,D{rr}/B{rr},"")', BLACK, NUM6)
        for k in range(W - i):
            tok = f"SUMIFS(Data_Cohort!$C:$C,Data_Cohort!$A:$A,$A{rr},Data_Cohort!$B:$B,{k})"
            put(FS, f"{TC(k)}{rr}", f"={tok}/{safe_days(f'EDATE($A{rr},{k})')}/1E9", BLACK, NUM3)
    ACT = lambda i, k: f"{TC(k)}{A0 + 1 + i}"  # noqa: E731
    R0 = A0 + W + 3
    put(FS, f"A{R0 - 1}", "Month-to-month development ratios, by cohort", BOLD)
    header(FS, R0, ["Cohort", "", ""] + [f"{k} to {k + 1}" for k in range(W - 1)])
    for i in range(W - 1):
        rr = R0 + 1 + i
        put(FS, f"A{rr}", f"=A{A0 + 1 + i}", BLACK, MONF)
        for k in range(W - 1 - i):
            put(FS, f"{TC(k)}{rr}", f'=IF({ACT(i, k)}>0,{ACT(i, k + 1)}/{ACT(i, k)},"")', BLACK, MULT)


    # inputs that reference the data
    INPUT_CELLS.update({"tail_factor.low": ("Forecast", f"B{r_tail}"), "tail_factor.base": ("Forecast", f"C{r_tail}"),
                        "tail_factor.high": ("Forecast", f"D{r_tail}"), "gateway_growth_decay": ("Forecast", f"C{r_decay}"),
                        "ent_usage_ratio.low": ("Forecast", f"B{r_ratio}"), "ent_usage_ratio.base": ("Forecast", f"C{r_ratio}"),
                        "ent_usage_ratio.high": ("Forecast", f"D{r_ratio}"), "ent_growth.low": ("Forecast", f"B{r_eg}"),
                        "ent_growth.base": ("Forecast", f"C{r_eg}"), "ent_growth.high": ("Forecast", f"D{r_eg}"),
                        "ent_ramp_days": ("Forecast", f"C{r_ramp}")})
    for s, v in zip(SCN, (0.93, 0.97, 1.0)):
        put(FS, f"{SC[s]}{r_tail}", v, BLUE, MULT, YEL if s == "base" else None)
    sig = lambda i: f"B{A0 + 1 + i}"  # noqa: E731
    put(FS, f"B{r_sg}", 0, BLUE, PCT)
    put(FS, f"C{r_sg}", f"=IF(AND({sig(T - 3)}>0,{sig(T)}>0),({sig(T)}/{sig(T - 3)})^(1/3)-1,0)", BLACK, PCT)
    put(FS, f"D{r_sg}", f"=2*C{r_sg}", BLACK, PCT)
    m0r = f"C{A0 + 1 + T - 2}:C{A0 + 1 + T}"
    put(FS, f"B{r_m0}", f"=IF(COUNT({m0r})=0,0,MIN({m0r}))", BLACK, NUM6)
    put(FS, f"C{r_m0}", f"=IF(COUNT({m0r})=0,0,AVERAGE({m0r}))", BLACK, NUM6)
    put(FS, f"D{r_m0}", f"=IF(COUNT({m0r})=0,0,MAX({m0r}))", BLACK, NUM6)
    put(FS, f"C{r_decay}", 0.5, BLUE, MULT, YEL)
    for s, v in zip(SCN, (0.8, 1.0, 1.2)):
        put(FS, f"{SC[s]}{r_ratio}", v, BLUE, PCT, YEL if s == "base" else None)
    for s, v in zip(SCN, (0.0, 0.03, 0.06)):
        put(FS, f"{SC[s]}{r_eg}", v, BLUE, PCT)
    put(FS, f"C{r_ramp}", 60, BLUE, NUM)
    asof_col = CL(2 + T)  # column of the as-of month on Monthly
    put(FS, f"C{r_ssp}", f"=IF(Monthly!{asof_col}{tok_rows[1]}=0,0,Monthly!{asof_col}{rev_rows[1]}/Monthly!{asof_col}{tok_rows[1]})", GREEN, USD4)
    put(FS, f"C{r_gap}", f"=IF(Monthly!{asof_col}{tok_rows[2]}=0,0,Monthly!{asof_col}{rev_rows[2]}/Monthly!{asof_col}{tok_rows[2]})", GREEN, USD4)
    put(FS, f"C{r_early}", f"=SUMIFS(Data_EarlierCohorts!$B:$B,Data_EarlierCohorts!$A:$A,{ASOF})/{safe_days(ASOF)}/1E9", BLACK, NUM3)

    # scenario projections
    row = R0 + W + 2
    NC = W + 3
    SCN_ROW = {}
    for s in SCN:
        put(FS, f"A{row}", f"Projected cohorts, {s} case (B tokens per day)", BOLD)
        hdr = row + 1
        header(FS, hdr, ["Cohort", "Signups", ""] + [f"Month {k}" for k in range(NC)])
        frow = hdr + 1 + NC
        for i in range(NC):
            rr = hdr + 1 + i
            put(FS, f"A{rr}", f"={mon(i)}", BLACK, MONF)
            put(FS, f"B{rr}", f"={sig(i)}" if i <= T else f"=B{rr - 1}*(1+${SC[s]}${r_sg})", BLACK, NUM)
            for k in range(NC - i):
                ref = f"{TC(k)}{rr}"
                if i + k <= T:
                    put(FS, ref, f"={ACT(i, k)}", BLACK, NUM3)
                elif k == 0:
                    put(FS, ref, f"=B{rr}*${SC[s]}${r_m0}", BLACK, NUM3)
                else:
                    put(FS, ref, f"={TC(k - 1)}{rr}*{TC(k - 1)}${frow}", BLACK, NUM3)
        put(FS, f"A{frow}", "Development factor, month k to k+1", BOLD)
        for k in range(NC - 1):
            tail = f"${SC[s]}${r_tail}"
            if k < T:
                top, bot = A0 + 1, A0 + 1 + (T - 1 - k)
                cur = f"{TC(k)}{top}:{TC(k)}{bot}"
                nxt = f"{TC(k + 1)}{top}:{TC(k + 1)}{bot}"
                if s == "base":
                    f = f"=IF(SUM({cur})=0,{tail},SUMPRODUCT(({cur}>0)*{nxt})/SUM({cur}))"
                else:
                    rng = f"{TC(k)}{R0 + 1}:{TC(k)}{R0 + 1 + (T - 1 - k)}"
                    f = f"=IF(COUNT({rng})=0,{tail},{'MIN' if s == 'low' else 'MAX'}({rng}))"
            else:
                f = f"={tail}"
            put(FS, f"{TC(k)}{frow}", f, BLACK, MULT)
        SCN_ROW[s] = (hdr, frow)
        row = frow + 3

    put(FS, f"A{row}", "Self-serve tokens per day by forecast month (B): sum of each diagonal + older cohorts", BOLD)
    cal = row + 1
    header(FS, cal, ["Scenario", f"=EDATE({ASOF},1)", f"=EDATE({ASOF},2)", f"=EDATE({ASOF},3)"])
    SS = {}
    for j, s in enumerate(SCN):
        rr = cal + 1 + j
        put(FS, f"A{rr}", s.capitalize())
        hdr, _ = SCN_ROW[s]
        for n in (1, 2, 3):
            mi = T + n
            cells = ",".join(f"{TC(mi - i)}{hdr + 1 + i}" for i in range(mi + 1))
            put(FS, f"{CL(1 + n)}{rr}", f"=SUM({cells})+$C${r_early}*${SC[s]}${r_tail}^{n}", BLACK, NUM2)
            SS[(s, n)] = f"{CL(1 + n)}{rr}"
    row = cal + 5

    put(FS, f"A{row}", "Gateways that don't pass customer IDs: tokens per day (B)", BOLD)
    gh = row + 1
    header(FS, gh, ["Line"] + [f"={mon(j)}" for j in range(W)])
    put(FS, f"A{gh + 1}", "Actual")
    for j in range(W):
        put(FS, f"{CL(2 + j)}{gh + 1}", f"=Monthly!{CL(2 + j)}{seg_tpd_rows['gateway_aggregate']}", GREEN, NUM2)
    gl, gp = f"{CL(2 + T)}{gh + 1}", f"{CL(1 + T)}{gh + 1}"
    put(FS, f"A{gh + 2}", "Last month's growth")
    put(FS, f"B{gh + 2}", f"=IF({gp}>0,{gl}/{gp}-1,0)", BLACK, PCT)
    put(FS, f"B{r_gw}", f"=MIN(0,B{gh + 2})", BLACK, PCT)
    put(FS, f"C{r_gw}", f"=B{gh + 2}*C{r_decay}", BLACK, PCT)
    put(FS, f"D{r_gw}", f"=MAX(0,B{gh + 2})", BLACK, PCT)
    GA = {(s, n): f"{gl}*(1+${SC[s]}${r_gw})^{n}" for s in SCN for n in (1, 2, 3)}
    row = gh + 5

    # summary
    put(FS, f"A{row}", "FORECAST SUMMARY", RED_B)
    sh = row + 1
    cols9 = [(s, n) for s in SCN for n in (1, 2, 3)]
    header(FS, sh, ["Tokens per day (B)"] + [f'="{s.capitalize()} "&TEXT(EDATE({ASOF},{n}),"mmm")' for s, n in cols9])
    comp = [("Self-serve (cohorts + older)", lambda s, n: f"={SS[(s, n)]}"),
            ("Gateways without customer IDs", lambda s, n: f"={GA[(s, n)]}"),
            ("Enterprise, signed contracts", lambda s, n: f"=Forecast_Ent!{CL(10 + SCN.index(s) * 3 + n)}{13 + MAX_CONTRACTS}"),
            ("Enterprise, pipeline", lambda s, n: f"=Forecast_Ent!{CL(1 + SCN.index(s) * 3 + n)}{{PIPE_TOK_ROW}}")]
    comp_rows = []
    for ci, (lab, fn) in enumerate(comp):
        rr = sh + 1 + ci
        put(FS, f"A{rr}", lab)
        for j, (s, n) in enumerate(cols9):
            put(FS, f"{CL(2 + j)}{rr}", fn(s, n), GREEN if "Forecast_Ent" in fn(s, n) else BLACK, NUM2)
        comp_rows.append(rr)
    TOT = sh + 1 + len(comp)
    put(FS, f"A{TOT}", "Total tokens per day (B)", BOLD)
    for j in range(9):
        put(FS, f"{CL(2 + j)}{TOT}", f"=SUM({CL(2 + j)}{comp_rows[0]}:{CL(2 + j)}{comp_rows[-1]})", BOLD, NUM2)
    put(FS, f"A{TOT + 1}", "Growth vs latest month actual")
    for j in range(9):
        put(FS, f"{CL(2 + j)}{TOT + 1}", f"=IF(Monthly!{asof_col}{M_TPD}=0,0,{CL(2 + j)}{TOT}/(Monthly!{asof_col}{M_TPD}-Monthly!{asof_col}{tok_rows[3]}/MAX(1,Monthly!{asof_col}{M_DAYS}))-1)", GREEN, PCT)
    rh = TOT + 3
    header(FS, rh, ["Net revenue ($k)"] + [f'="{s.capitalize()} "&TEXT(EDATE({ASOF},{n}),"mmm")' for s, n in cols9])
    days_n = lambda n: f"DAY(EOMONTH({ASOF},{n}))"  # noqa: E731
    rev = [("Self-serve", lambda s, n, c: f"={c}{comp_rows[0]}*{days_n(n)}*1000*$C${r_ssp}/1000"),
           ("Gateways without customer IDs", lambda s, n, c: f"={c}{comp_rows[1]}*{days_n(n)}*1000*$C${r_gap}/1000"),
           ("Enterprise, signed", lambda s, n, c: f"=Forecast_Ent!{CL(1 + SCN.index(s) * 3 + n)}{{ENT_REV_ROW}}/1000"),
           ("Enterprise, pipeline", lambda s, n, c: f"=Forecast_Ent!{CL(1 + SCN.index(s) * 3 + n)}{{PIPE_REV_ROW}}/1000")]
    rev_rows_f = []
    for li, (lab, fn) in enumerate(rev):
        rr = rh + 1 + li
        put(FS, f"A{rr}", lab)
        for j, (s, n) in enumerate(cols9):
            put(FS, f"{CL(2 + j)}{rr}", fn(s, n, CL(2 + j)), BLACK, NUM1)
        rev_rows_f.append(rr)
    REV = rh + 1 + len(rev)
    put(FS, f"A{REV}", "Total net revenue ($k)", BOLD)
    for j in range(9):
        put(FS, f"{CL(2 + j)}{REV}", f"=SUM({CL(2 + j)}{rev_rows_f[0]}:{CL(2 + j)}{rev_rows_f[-1]})", BOLD, NUM1)

    # variance against last close's stored forecast
    vr = REV + 3
    put(FS, f"A{vr}", "How good was last month's forecast? Latest month actual vs the base forecast stored at the previous close (Data_FcHistory)", BOLD)
    header(FS, vr + 1, ["Tokens per day (B)", "Forecast", "Actual", "Variance", "Variance %", "Flag (>10%)"])
    prev_close = f"EDATE({ASOF},-1)"
    comps = [("self_serve", "Self-serve", f"Monthly!{asof_col}{seg_tpd_rows['self_serve']}"),
             ("gateway_aggregate", "Gateways without customer IDs", f"Monthly!{asof_col}{seg_tpd_rows['gateway_aggregate']}"),
             ("enterprise", "Enterprise (signed + pipeline)", f"Monthly!{asof_col}{seg_tpd_rows['enterprise']}")]
    v0 = vr + 2
    for i, (key, lab, act) in enumerate(comps):
        rr = v0 + i
        crit = f'Data_FcHistory!$A:$A,{prev_close},Data_FcHistory!$B:$B,{ASOF},Data_FcHistory!$C:$C,"base",Data_FcHistory!$D:$D,"{key}"'
        put(FS, f"A{rr}", lab)
        put(FS, f"B{rr}", f'=IF(COUNTIFS({crit})=0,"none on file",SUMIFS(Data_FcHistory!$E:$E,{crit}))', BLACK, NUM2)
        put(FS, f"C{rr}", f"={act}", GREEN, NUM2)
        put(FS, f"D{rr}", f'=IF(ISNUMBER(B{rr}),C{rr}-B{rr},"")', BLACK, NUM2)
        put(FS, f"E{rr}", f'=IF(AND(ISNUMBER(B{rr}),N(B{rr})<>0),D{rr}/B{rr},"")', BLACK, PCT)
        put(FS, f"F{rr}", f'=IF(E{rr}="","",IF(ABS(E{rr})>0.1,"FLAG","ok"))', BLACK)
    VT = v0 + len(comps)
    put(FS, f"A{VT}", "Total", BOLD)
    put(FS, f"B{VT}", f'=IF(COUNT(B{v0}:B{VT - 1})=0,"none on file",SUM(B{v0}:B{VT - 1}))', BOLD, NUM2)
    put(FS, f"C{VT}", f"=SUM(C{v0}:C{VT - 1})", BOLD, NUM2)
    put(FS, f"D{VT}", f'=IF(ISNUMBER(B{VT}),C{VT}-B{VT},"")', BOLD, NUM2)
    put(FS, f"E{VT}", f'=IF(AND(ISNUMBER(B{VT}),N(B{VT})<>0),D{VT}/B{VT},"")', BOLD, PCT)
    put(FS, f"F{VT}", f'=IF(E{VT}="","",IF(ABS(E{VT})>0.1,"FLAG","ok"))', BOLD)
    put(FS, f"A{VT + 1}", "The runner stores each close's forecast in forecast_history.csv. On the first close it reconstructs the previous close's forecast from data (without pipeline, since no old CRM snapshot exists).", SUB)
    widths(FS, {"A": 50, "B": 12, "C": 12, "D": 12, "E": 12, **{CL(j): 10 for j in range(6, 4 + NC + 1)}})
    FS.freeze_panes = "B4"

    # ======================================================================== Forecast_Ent
    FE = wb.create_sheet("Forecast_Ent")
    title(FE, "Enterprise forecast", "Contracts: last-7-day run-rate ramping toward commitment x usage ratio for young contracts, then monthly growth. Pipeline: probability-weighted with the same ramp.")
    put(FE, "A4", "Month helpers", BOLD)
    for i, lab in enumerate(["Forecast month", "First day", "Last day", "Days", "Months after close"]):
        put(FE, f"A{5 + i}", lab)
    for n in (1, 2, 3):
        col = CL(1 + n)
        put(FE, f"{col}5", f"=EDATE({ASOF},{n})", BOLD, MONF)
        put(FE, f"{col}6", f"=EDATE({ASOF},{n})", BLACK, DATEF)
        put(FE, f"{col}7", f"=EOMONTH({col}6,0)", BLACK, DATEF)
        put(FE, f"{col}8", f"={col}7-{col}6+1", BLACK, NUM)
        put(FE, f"{col}9", n, BLACK, NUM)
    put(FE, "F5", "Close date", BOLD)
    put(FE, "F6", f"=EOMONTH({ASOF},0)", BLACK, DATEF)
    put(FE, "F7", "List price ($/M)", BOLD)
    put(FE, "F8", f"={P('list_price_per_mtok')}", GREEN, USD4)
    FI = lambda row, s: f"Forecast!${SC[s]}${row}"  # noqa: E731
    RAMP = f"Forecast!$C${r_ramp}"
    r = 12
    put(FE, f"A{r - 1}", "Signed contracts: tokens per day (B)", BOLD)
    cols = ["Customer", "Start", "End", "Monthly commit ($)", "Price ($/M)", "Age at close (days)", "Last 7 days, tokens/day (B)",
            "Steady state, low", "Steady state, base", "Steady state, high"] + [f"{s.capitalize()} M+{n}" for s in SCN for n in (1, 2, 3)]
    header(FE, r, cols)
    c0 = r + 1
    for n_ in range(1, MAX_CONTRACTS + 1):
        rr = c0 + n_ - 1
        put(FE, f"A{rr}", f"={ix('Data_Contracts', 'A', n_)}")
        put(FE, f"B{rr}", f"={ix('Data_Contracts', 'C', n_)}", BLACK, DATEF)
        put(FE, f"C{rr}", f"={ix('Data_Contracts', 'D', n_)}", BLACK, DATEF)
        put(FE, f"D{rr}", f"={ix('Data_Contracts', 'E', n_)}", BLACK, USD)
        put(FE, f"E{rr}", f"={ix('Data_Contracts', 'F', n_)}", BLACK, USD4)
        put(FE, f"F{rr}", f'=IF($A{rr}="","",$F$6-$B{rr})', BLACK, NUM)
        put(FE, f"G{rr}", f'=IF($A{rr}="",0,N(INDEX(Data_Contracts!$H:$H,{n_}+1)))', BLACK, NUM3)
        for j, s in enumerate(SCN):
            put(FE, f"{CL(8 + j)}{rr}", f'=IF(OR($A{rr}="",N($E{rr})=0),0,$D{rr}/$E{rr}*{FI(r_ratio, s)}/30.4/1000)', BLACK, NUM3)
        for j, s in enumerate(SCN):
            st = CL(8 + j)
            for n in (1, 2, 3):
                col = CL(11 + j * 3 + n - 1)
                mc = CL(1 + n)
                put(FE, f"{col}{rr}",
                    f'=IF(OR($A{rr}="",AND($C{rr}<>"",N($C{rr})<{mc}$6)),0,'
                    f'IF($F{rr}<{RAMP},$G{rr}+({st}{rr}-$G{rr})*MIN(1,{mc}$9*30.4/({RAMP}-$F{rr})),$G{rr})*(1+{FI(r_eg, s)})^{mc}$9)',
                    BLACK, NUM3)
    c_last = c0 + MAX_CONTRACTS - 1
    CT = c_last + 1
    put(FE, f"A{CT}", "Total", BOLD)
    for j in range(9):
        col = CL(11 + j)
        put(FE, f"{col}{CT}", f"=SUM({col}{c0}:{col}{c_last})", BOLD, NUM2)
    assert CT == 13 + MAX_CONTRACTS  # the Forecast summary points at this row

    r = CT + 3
    put(FE, f"A{r - 1}", "Signed contracts: revenue ($) = greater of usage x price and the commitment", BOLD)
    header(FE, r, ["Customer"] + [f"{s.capitalize()} M+{n}" for s in SCN for n in (1, 2, 3)])
    rv0 = r + 1
    for n_ in range(MAX_CONTRACTS):
        rr = rv0 + n_
        tr = c0 + n_
        put(FE, f"A{rr}", f"=A{tr}")
        for j, s in enumerate(SCN):
            for n in (1, 2, 3):
                col = CL(1 + j * 3 + n)
                tcol = CL(11 + j * 3 + n - 1)
                mc = CL(1 + n)
                put(FE, f"{col}{rr}", f'=IF(OR($A{tr}="",AND($C{tr}<>"",N($C{tr})<{mc}$6)),0,MAX({tcol}{tr}*{mc}$8*1000*$E{tr},$D{tr}))', BLACK, USD)
    RT = rv0 + MAX_CONTRACTS
    put(FE, f"A{RT}", "Total", BOLD)
    for j in range(9):
        col = CL(2 + j)
        put(FE, f"{col}{RT}", f"=SUM({col}{rv0}:{col}{RT - 1})", BOLD, USD)
    ENT_REV_ROW = RT

    r = RT + 3
    put(FE, f"A{r - 1}", "Open pipeline: low = none, base = probability-weighted, high = deals at 50%+ counted in full", BOLD)
    pc = ["Deal", "Customer", "Probability", "Expected commit ($)", "Expected start", "Discount", "Price ($/M)",
          "Steady tokens/day (B)", "Weight low", "Weight base", "Weight high"]
    for n in (1, 2, 3):
        pc += [f"Active days M+{n}", f"Avg ramp M+{n}"]
    header(FE, r, pc)
    q0 = r + 1
    for n_ in range(1, MAX_DEALS + 1):
        rr = q0 + n_ - 1
        put(FE, f"A{rr}", f"={ix('Data_Pipeline', 'A', n_)}")
        put(FE, f"B{rr}", f"={ix('Data_Pipeline', 'B', n_)}")
        put(FE, f"C{rr}", f'=IF($A{rr}="",0,N(INDEX(Data_Pipeline!$D:$D,{n_}+1)))', BLACK, PCT)
        put(FE, f"D{rr}", f'=IF($A{rr}="",0,N(INDEX(Data_Pipeline!$E:$E,{n_}+1)))', BLACK, USD)
        put(FE, f"E{rr}", f"={ix('Data_Pipeline', 'F', n_)}", BLACK, DATEF)
        put(FE, f"F{rr}", f'=IF($A{rr}="",0,N(INDEX(Data_Pipeline!$G:$G,{n_}+1)))', BLACK, PCT)
        put(FE, f"G{rr}", f"=$F$8*(1-F{rr})", BLACK, USD4)
        put(FE, f"H{rr}", f"=IF(G{rr}=0,0,D{rr}/G{rr}/30.4/1000)", BLACK, NUM3)
        put(FE, f"I{rr}", 0, BLACK, PCT)
        put(FE, f"J{rr}", f"=C{rr}", BLACK, PCT)
        put(FE, f"K{rr}", f"=IF(C{rr}>=0.5,1,C{rr})", BLACK, PCT)
        for n in (1, 2, 3):
            mc = CL(1 + n)
            ad, rp = CL(12 + (n - 1) * 2), CL(13 + (n - 1) * 2)
            put(FE, f"{ad}{rr}", f'=IF(OR($A{rr}="",$E{rr}=""),0,MAX(0,{mc}$7-MAX($E{rr},{mc}$6)+1))', BLACK, NUM)
            put(FE, f"{rp}{rr}", f'=IF({ad}{rr}=0,0,MIN(1,0.1+0.9*((MAX($E{rr},{mc}$6)+{mc}$7)/2-$E{rr})/{RAMP}))', BLACK, PCT)
    q_last = q0 + MAX_DEALS - 1
    r = q_last + 3
    put(FE, f"A{r - 1}", "Open pipeline: tokens per day (B) and revenue ($)", BOLD)
    header(FE, r, ["Line"] + [f"{s.capitalize()} M+{n}" for s in SCN for n in (1, 2, 3)])
    put(FE, f"A{r + 1}", "Tokens per day (B)")
    put(FE, f"A{r + 2}", "Revenue ($)")
    rng = lambda c: f"{c}{q0}:{c}{q_last}"  # noqa: E731
    for j, (s, wcol) in enumerate(zip(SCN, "IJK")):
        for n in (1, 2, 3):
            col = CL(1 + j * 3 + n)
            ad, rp = CL(12 + (n - 1) * 2), CL(13 + (n - 1) * 2)
            put(FE, f"{col}{r + 1}", f"=SUMPRODUCT({rng(wcol)},{rng('H')},{rng(ad)},{rng(rp)})/{CL(1 + n)}$8", BLACK, NUM3)
            put(FE, f"{col}{r + 2}", f"=SUMPRODUCT({rng(wcol)},{rng('H')},{rng(ad)},{rng(rp)},{rng('G')})*1000", BLACK, USD)
    PIPE_TOK_ROW, PIPE_REV_ROW = r + 1, r + 2
    widths(FE, {"A": 24, "B": 20, **{CL(j): 12 for j in range(3, 20)}})
    FE.freeze_panes = "B5"

    # resolve the forward references in Forecast
    for row_ in FS.iter_rows():
        for c in row_:
            if isinstance(c.value, str) and "{" in c.value and "_ROW}" in c.value:
                c.value = (c.value.replace("{PIPE_TOK_ROW}", str(PIPE_TOK_ROW)).replace("{ENT_REV_ROW}", str(ENT_REV_ROW))
                           .replace("{PIPE_REV_ROW}", str(PIPE_REV_ROW)))

    FC_TOT = {(s, n): f"Forecast!{CL(2 + SCN.index(s) * 3 + n - 1)}{TOT}" for s in SCN for n in (1, 2, 3)}
    FC_REV = {(s, n): f"Forecast!{CL(2 + SCN.index(s) * 3 + n - 1)}{REV}" for s in SCN for n in (1, 2, 3)}

    # ======================================================================== Capacity
    CP = wb.create_sheet("Capacity")
    title(CP, "From demand forecast to GPUs: reserve the base, rent the peaks", "Inference serving only. Peak and base ratios come from Data_Hourly when hourly data exists; otherwise the override inputs are used.")
    put(CP, "A4", "Inputs", BOLD)
    header(CP, 5, ["Input", "Value", "Source / why"])
    HB = "Data_Hourly!$B$2:$B$169"
    cin = [("Tokens per second per busy GPU", 20604, NUM, BLUE, "Replace with measured serving throughput. Placeholder from the outside-in model: H100, BF16, 8B params."),
           ("Target utilization at peak", 0.85, PCT, BLUE, "Headroom for the latency promise and bursts."),
           ("Reserved rate ($ / GPU-hour)", 2.50, USD2, BLUE, "From your capacity contract."),
           ("On-demand rate ($ / GPU-hour)", 3.00, USD2, BLUE, "Same provider tier."),
           ("Peak hour / average (override if no hourly data)", 1.75, MULT, BLUE, "Used only when Data_Hourly is empty."),
           ("Quietest hour / average (override if no hourly data)", 0.55, MULT, BLUE, "Used only when Data_Hourly is empty."),
           ("Peak hour / average (used)", f"=IF(COUNT({HB})=0,B10,MAX({HB})/AVERAGE({HB}))", MULT, BLACK, "Last 7 days of the as-of month."),
           ("Quietest hour / average (used)", f"=IF(COUNT({HB})=0,B11,MIN({HB})/AVERAGE({HB}))", MULT, BLACK, ""),
           ("Last 7 days: tokens per day (B)", f'=IF(COUNT({HB})=0,"",SUM({HB})/7/1E9)', NUM2, BLACK, "")]
    for i, (lab, v, fmt, font, why) in enumerate(cin):
        rr = 6 + i
        put(CP, f"A{rr}", lab)
        put(CP, f"B{rr}", v, font, fmt, YEL if i in (0, 1) else None)
        put(CP, f"C{rr}", why, SUB)
    for k, rr in (("gpu_tokens_per_sec", 6), ("peak_utilization", 7), ("reserved_rate", 8), ("on_demand_rate", 9),
                  ("peak_ratio_override", 10), ("base_ratio_override", 11)):
        INPUT_CELLS[f"capacity.{k}"] = ("Capacity", f"B{rr}")
    TPS, UTIL, RRATE, ORATE, PEAKR, BASER, TPD7 = "$B$6", "$B$7", "$B$8", "$B$9", "$B$12", "$B$13", "$B$14"
    r = 17
    put(CP, f"A{r - 1}", "GPUs needed by month and scenario", BOLD)
    header(CP, r, ["Line"] + [f'="{s.capitalize()} "&TEXT(EDATE({ASOF},{n}),"mmm")' for s, n in cols9])
    lines = ["Tokens per day (B)", "Average tokens/sec", "Peak-hour tokens/sec", "Base-load tokens/sec", "GPUs at peak", "GPUs for base load"]
    for li, lab in enumerate(lines):
        put(CP, f"A{r + 1 + li}", lab)
    for j, (s, n) in enumerate(cols9):
        col = CL(2 + j)
        put(CP, f"{col}{r + 1}", f"={FC_TOT[(s, n)]}", GREEN, NUM2)
        put(CP, f"{col}{r + 2}", f"={col}{r + 1}*1E9/86400", BLACK, NUM)
        put(CP, f"{col}{r + 3}", f"={col}{r + 2}*{PEAKR}", BLACK, NUM)
        put(CP, f"{col}{r + 4}", f"={col}{r + 2}*{BASER}", BLACK, NUM)
        put(CP, f"{col}{r + 5}", f"=ROUNDUP({col}{r + 3}/{TPS}/{UTIL},0)", BOLD, NUM)
        put(CP, f"{col}{r + 6}", f"=ROUNDDOWN({col}{r + 4}/{TPS},0)", BOLD, NUM)
    G_PEAK, G_BASE = r + 5, r + 6
    r = G_BASE + 3
    put(CP, f"A{r}", "Recommendation", RED_B)
    put(CP, f"A{r + 1}", "Reserve: base-load GPUs in the LOW case, next month (capacity that stays busy)")
    put(CP, f"B{r + 1}", f"=B{G_BASE}", BOLD, NUM)
    RES = f"$B${r + 1}"
    put(CP, f"A{r + 2}", "Base case peak on top of that, rented on demand")
    put(CP, f"B{r + 2}", f"=MAX(0,E{G_PEAK}-{RES})", BOLD, NUM)
    put(CP, f"A{r + 3}", "High case peak: on-demand headroom to line up with the provider")
    put(CP, f"B{r + 3}", f"=MAX(0,H{G_PEAK}-{RES})", BOLD, NUM)
    REC_ROW = r + 1
    r += 6
    put(CP, f"A{r - 1}", "Monthly serving cost, base case: three ways to buy the same capacity", BOLD)
    header(CP, r, ["Strategy", f"=EDATE({ASOF},1)", f"=EDATE({ASOF},2)", f"=EDATE({ASOF},3)", "How it's calculated"])
    for si, (lab, kind) in enumerate([("Reserve the base, rent the peaks", "mix"), ("Reserve enough for the peak", "peak"), ("Rent everything on demand", "od")]):
        rr = r + 1 + si
        put(CP, f"A{rr}", lab, BOLD if kind == "mix" else BLACK)
        for n in (1, 2, 3):
            bcol = CL(4 + n)  # base-case columns E,F,G
            days = f"DAY(EOMONTH({ASOF},{n}))"
            need = f"({HB}/3600*({bcol}{G_PEAK - 4}/{TPD7})/{TPS}/{UTIL})"
            if kind == "mix":
                f = f'=IF(COUNT({HB})=0,"needs hourly data",{RES}*24*{days}*{RRATE}+SUMPRODUCT(({need}-{RES})*({need}>{RES}))/7*{days}*{ORATE})'
            elif kind == "peak":
                f = f"={bcol}{G_PEAK}*24*{days}*{RRATE}"
            else:
                f = f'=IF(COUNT({HB})=0,"needs hourly data",SUMPRODUCT({need})/7*{days}*{ORATE})'
            put(CP, f"{CL(1 + n)}{rr}", f, BLACK, USD)
        put(CP, f"E{rr}", {"mix": "Reserved GPUs x all hours x reserved rate + GPU-hours above that each hour x on-demand",
                           "peak": "Peak GPUs x all hours x reserved rate (idle off-peak)",
                           "od": "GPU-hours needed each hour x on-demand rate"}[kind], SUB)
    put(CP, f"A{r + 4}", "Base-case net revenue ($)", BOLD)
    put(CP, f"A{r + 5}", "Serving gross margin, recommended strategy", BOLD)
    for n in (1, 2, 3):
        col = CL(1 + n)
        put(CP, f"{col}{r + 4}", f"={FC_REV[('base', n)]}*1000", GREEN, USD)
        put(CP, f"{col}{r + 5}", f'=IF(OR(NOT(ISNUMBER({col}{r + 1})),{col}{r + 4}=0),"",1-{col}{r + 1}/{col}{r + 4})', BOLD, PCTR)
    put(CP, f"A{r + 7}", "Inference compute only: excludes training, networking, storage and on-call staff. Hourly need = the last 7 days' hourly shape scaled to each month's volume.", SUB)
    widths(CP, {"A": 58, "B": 13, "C": 13, "D": 13, **{CL(j): 12 for j in range(5, 11)}})

    # ======================================================================== Summary
    S = wb.create_sheet("Summary", 0)
    title(S, "Monthly close", "Formulas only. Start here, then Checks. Every figure traces to a Data_ sheet and from there to a query in /sql.")
    put(S, "A3", "As-of month", BOLD); put(S, "B3", f"={ASOF}", GREEN, MONF)
    put(S, "C3", "Sources", BOLD); put(S, "D3", f"={P('sources_loaded')}", GREEN)
    six = [CL(2 + j) for j in range(WINDOW - 6, WINDOW)]
    header(S, 5, ["Metric"] + [f"=Monthly!{c}4" for c in six])
    kpi = [("Billable tokens per day (B)", M_TPD, NUM1), ("Growth in tokens per day", M_GR, PCT), ("Net revenue ($k)", M_REV, NUM1),
           ("Annualized run-rate ($k)", M_ARR, NUM), ("Enterprise share of revenue", M_ENTSH, PCT),
           ("Realized net price per M tokens", M_PRICE, USD4), ("Active accounts", M_ACT, NUM), ("Paying accounts", M_PAY, NUM),
           ("Failed calls (not billed)", M_FAIL, PCT)]
    for i, (lab, mr, fmt) in enumerate(kpi):
        rr = 6 + i
        put(S, f"A{rr}", lab)
        for j, c in enumerate(six):
            put(S, f"{CL(2 + j)}{rr}", f"=Monthly!{c}{mr}", GREEN, fmt)
    r = 6 + len(kpi) + 1
    put(S, f"A{r}", "Close checks", BOLD)
    G = f"Checks!$G$5:$G${4 + MAX_CHECK_ROWS}"
    for i, (lab, st) in enumerate([("Flagged for review", "FLAG"), ("Informational", "INFO"), ("Skipped (source not connected)", "SKIP"), ("Passed", "PASS")]):
        put(S, f"A{r + 1 + i}", lab)
        put(S, f"B{r + 1 + i}", f'=COUNTIF({G},"{st}")', GREEN, NUM)
    put(S, f"A{r + 5}", "  $ impact of flagged items")
    put(S, f"B{r + 5}", f'=SUMIF({G},"FLAG",Checks!$F$5:$F${4 + MAX_CHECK_ROWS})', GREEN, USD2)
    r += 7
    put(S, f"A{r}", "Forecast, tokens per day (B)", BOLD)
    header(S, r + 1, ["Scenario", f"=EDATE({ASOF},1)", f"=EDATE({ASOF},2)", f"=EDATE({ASOF},3)"])
    for j, s in enumerate(SCN):
        put(S, f"A{r + 2 + j}", s.capitalize())
        for n in (1, 2, 3):
            put(S, f"{CL(1 + n)}{r + 2 + j}", f"={FC_TOT[(s, n)]}", GREEN, NUM1)
    put(S, f"A{r + 5}", "Net revenue, base ($k)")
    for n in (1, 2, 3):
        put(S, f"{CL(1 + n)}{r + 5}", f"={FC_REV[('base', n)]}", GREEN, NUM1)
    put(S, f"A{r + 6}", "Last month's forecast vs actual (total)")
    put(S, f"B{r + 6}", f'=IF(Forecast!E{VT}="","no forecast on file",Forecast!E{VT})', GREEN, PCT)
    put(S, f"A{r + 7}", "GPUs to reserve for serving")
    put(S, f"B{r + 7}", f"=Capacity!B{REC_ROW}", GREEN, NUM)
    r += 9
    put(S, f"A{r}", "Commentary (written by finance, or drafted by the agent per AGENT_BRIEF.md and reviewed before it goes anywhere)", BOLD)
    for i in range(6):
        S.merge_cells(f"A{r + 1 + i}:G{r + 1 + i}")
        S[f"A{r + 1 + i}"].fill = INFILL
        S[f"A{r + 1 + i}"].alignment = WRAP
        S.row_dimensions[r + 1 + i].height = 30
    widths(S, {"A": 44, **{CL(j): 13 for j in range(2, 8)}})

    # ======================================================================== Start
    ST = wb.create_sheet("Start", 0)
    put(ST, "A1", "Monthly close workbook", TITLE)
    rows = [
        ("How this workbook works", None),
        ("1. Numbers enter only through the grey Data_ tabs. The runner fills them from your sources; you can also paste query results by hand.", None),
        ("2. Every other tab is formulas reading the Data_ tabs for a 12-month window ending at the as-of month.", None),
        ("3. Blue cells are judgment inputs (forecast and capacity). The runner seeds them from config.yaml; edit them here to test scenarios.", None),
        ("4. Grey-shaded cells on Checks and Summary are for people and the agent: owners, explanations, commentary.", None),
        ("", None),
        ("This close", None),
        ("As-of month", f"={ASOF}"),
        ("Generated", f"={P('generated_at')}"),
        ("Usage grain", f"={P('usage_grain')}"),
        ("Sources connected", f"={P('sources_loaded')}"),
        ("Usage data from", f"={P('data_first_date')}"),
        ("Usage data to", f"={P('data_last_date')}"),
        ("Checks flagged", f'=COUNTIF({G},"FLAG")'),
        ("", None),
        ("Color key: blue = input, black = formula, green = link to another tab, yellow = key assumption, grey tab = data.", None),
    ]
    fmts = {"As-of month": MONF, "Usage data from": DATEF, "Usage data to": DATEF}
    for i, (a, b) in enumerate(rows):
        rr = 3 + i
        put(ST, f"A{rr}", a, BOLD if a in ("How this workbook works", "This close") else (SUB if a.startswith("Color") else BLACK))
        if b:
            put(ST, f"B{rr}", b, GREEN, fmts.get(a))
    widths(ST, {"A": 40, "B": 30})

    order = ["Start", "Summary", "Checks", "Monthly", "UseCases", "Cohorts", "Enterprise", "Concentration",
             "Forecast", "Forecast_Ent", "Capacity"] + list(DATA_SHEETS)
    wb._sheets = [wb[n] for n in order]
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    wb.active = 0
    wb.save(path)
    import json
    from pathlib import Path
    Path(path).with_name("input_cells.json").write_text(json.dumps(INPUT_CELLS, indent=1))
    return INPUT_CELLS
