"""
Overrides tab: bad entries must never break the workbook, and the switch must restore it exactly.

    python tests/test_overrides.py      (needs LibreOffice to recalculate; run_tests.py calls it)

1. Baseline: the example close with no overrides.
2. Every kind of bad entry (text, an error, out of range, wrong month, unknown segment) plus a few
   valid ones, switch ON: no new error cells anywhere, valid ones applied, bad ones ignored.
3. The same entries with the switch OFF: every number on every report tab equals the baseline.
"""

import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "examples/outputs/Close_2027-02.xlsx"
TMP = ROOT / "tests/outputs/overrides"
OFFICE = shutil.which("soffice") or shutil.which("libreoffice")
REPORT = ["Start", "Summary", "Checks", "Monthly", "UseCases", "Cohorts", "Enterprise", "Concentration",
          "Forecast", "Forecast_Ent", "Capacity", "Overrides"]


def ok(cond, msg):
    print(("  PASS " if cond else "  FAIL ") + msg)
    if not cond:
        raise SystemExit(1)


def recalc(path):
    out = TMP / "calc"
    out.mkdir(parents=True, exist_ok=True)
    subprocess.run([OFFICE, "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(path)], capture_output=True, check=True)
    return load_workbook(out / path.name, data_only=True)


def values(wb):
    return {(ws.title, c.coordinate): c.value for ws in (wb[n] for n in REPORT) for row in ws.iter_rows() for c in row}


ERRS = {"#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:502", "Err:504", "Err:522"}


def errors(vals):
    """Error cells outside the override inputs themselves (column D on Overrides is where bad entries are typed)."""
    return sorted(k for k, v in vals.items() if v in ERRS
                  and not (k[0] == "Overrides" and (k[1].startswith("D") or (k[1][0] in "ABC" and int(k[1][1:]) >= 46))))


def row_of(ws, key):
    return next(r for r in range(8, 40) if ws.cell(r, 1).value == key)


def edit(switch, name):
    wb = load_workbook(SRC)
    ov = wb["Overrides"]
    ov["C4"] = switch
    bad = {"tail_factor.base": "abc", "ent_ramp_days": 0, "capacity.reserved_rate": -5,
           "capacity.gpu_tokens_per_sec": "=1/0", "capacity.peak_utilization": "=NA()"}
    good = {"capacity.on_demand_rate": 4.5, "price.self_serve": 0.05}
    for k, v in {**bad, **good}.items():
        if v is not None:
            ov.cell(row_of(ov, k), 4, v)
    adj = [(date(2027, 2, 15), "enterprise", -1458.67),      # valid
           (date(2020, 1, 1), "enterprise", 500),           # outside the window
           (date(2027, 1, 1), "nonsense", 500),             # unknown segment
           (date(2027, 1, 1), "self_serve", "lots"),        # amount not a number
           ("last month", "self_serve", 100),               # month not a date
           (date(2027, 1, 1), "self_serve", "=1/0")]        # error value
    for i, (m, s, a) in enumerate(adj):
        for j, v in enumerate((m, s, a)):
            ov.cell(46 + i, 1 + j, v)
    p = TMP / name
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def main():
    if not OFFICE:
        print("  SKIP LibreOffice not installed")
        return
    shutil.rmtree(TMP, ignore_errors=True)
    TMP.mkdir(parents=True)
    base_p = TMP / "base.xlsx"
    shutil.copy(SRC, base_p)
    base = values(recalc(base_p))
    ok(not errors(base), f"baseline has no error cells ({len(errors(base))})")

    on = recalc(edit("ON", "on.xlsx"))
    von = values(on)
    ok(not errors(von), f"bad overrides cause no error cells (found {errors(von)[:5]})")
    ov = on["Overrides"]
    st = {ov.cell(r, 1).value: ov.cell(r, 6).value for r in range(8, 40) if ov.cell(r, 1).value}
    ok(st["tail_factor.base"] == "Ignored: not a number", "text is ignored")
    ok(st["ent_ramp_days"] == "Ignored: outside range" and st["capacity.reserved_rate"] == "Ignored: outside range", "out-of-range numbers are ignored")
    ok(st["capacity.gpu_tokens_per_sec"] == "Ignored: error value" and st["capacity.peak_utilization"] == "Ignored: error value", "error values are ignored")
    ok(st["capacity.on_demand_rate"] == "Applied" and on["Capacity"]["B9"].value == 4.5, "valid override reaches Capacity")
    ok(abs(on["Capacity"]["B6"].value - base[("Capacity", "B6")]) < 1e-9, "ignored override keeps the model value")
    adj = [ov.cell(46 + i, 8).value for i in range(6)]
    ok(adj == ["Applied", "Ignored: month outside the window", "Ignored: unknown segment",
               "Ignored: amount is not a number", "Ignored: month is not a date", "Ignored: error value"], f"adjustment checks: {adj}")
    m = on["Monthly"]
    tot = next(r for r in range(1, 80) if m.cell(r, 1).value == "Total net revenue")
    feb = m.cell(tot, 13).value
    ok(abs(feb - (base[("Monthly", f"M{tot}")] - 1.45867)) < 1e-6, "valid adjustment flows into Feb net revenue")
    fs = on["Forecast"]
    ftot = next(r for r in range(1, fs.max_row + 1) if fs.cell(r, 1).value == "Total tokens per day (B)")
    ok(all(isinstance(fs.cell(ftot, c).value, (int, float)) for c in range(2, 11)), "forecast totals still calculate")
    ok("1 adjustment" in str(on["Overrides"]["C5"].value) and "2 input" in str(on["Overrides"]["C5"].value), f"banner counts: {on['Overrides']['C5'].value}")

    off = values(recalc(edit("OFF", "off.xlsx")))
    diffs = [k for k in base if k[0] != "Overrides" and k != ("Summary", "G3") and k != ("Start", "B17")
             and isinstance(base[k], (int, float)) and abs((off.get(k) or 0) - base[k]) > 1e-9 * max(1, abs(base[k]))]
    ok(not diffs, f"switch OFF restores every number on every report tab ({len(diffs)} differ: {diffs[:5]})")
    ok(not errors(off), "switch OFF has no error cells")


if __name__ == "__main__":
    sys.exit(main())
